import shutil
from pathlib import Path

import openpyxl
import pytest

from app.aggregate import to_monthly
from app.sample_data import make_month
from app.schema import load_template
from app.workbook import GNWorkbook, WorkbookError

T = load_template()
SRC = Path(__file__).parents[2] / "samples" / "gn_workbook_template.xlsx"


@pytest.fixture
def wbpath(tmp_path):
    p = tmp_path / "gn.xlsx"
    shutil.copy(SRC, p)
    return p


def test_monthly_rollup():
    m = to_monthly(make_month(T), T)
    assert m["Total number of meetings held"] == 4
    assert m["Number of SHG members"] == 15
    assert m["Total attendance for the month"] == 54
    assert m["Savings (Rs.)"] == 11400
    assert m["Loans distributed (Rs.)"] == 14000
    assert m["Loan Purpose – Income generation (Rs.)"] == 13000
    assert m["Loan right off (Rs.)"] == 0
    assert m["Number of members who attended the last meeting"] == 12
    # month-end mother-book cash = last week's cash in hand
    rec = make_month(T)
    assert m["Cash in Hand – at end of month (mother book) (Rs.)"] == rec.get("weekly.cash_in_hand.W4")


def test_new_shg_first_month_creates_tab_and_opening_balances(wbpath):
    gn = GNWorkbook(wbpath, T)
    rep = gn.write_month("Kalaimagal", "2026-10", to_monthly(make_month(T), T),
                         gn_name="Puthukkudiyiruppu East", create_tab=True)
    gn.save()
    assert rep.created_tab and rep.column == "C"
    assert "Number of members who attended the last meeting" in rep.skipped   # no such row yet
    wb = openpyxl.load_workbook(wbpath)
    ws = wb["Kalaimagal"]
    assert ws["C6"].value == 11400                          # Savings, month 1
    assert ws["C21"].value == 11400                         # opening savings to date
    assert ws["D21"].value == "=C21+D6"                     # formulas untouched
    mon = wb["SHG Monitoring"]
    names = [mon.cell(r, 2).value for r in range(2, mon.max_row + 1)]
    assert names.count("Kalaimagal") == 29                  # block registered


def test_second_month_goes_to_next_column_and_prior_is_read(wbpath):
    gn = GNWorkbook(wbpath, T)
    gn.write_month("Kalaimagal", "2026-10", to_monthly(make_month(T), T), gn_name="PTK East", create_tab=True)
    nov = make_month(T, month="2026-11")
    rep = gn.write_month("kalaimagal ", "2026-11", to_monthly(nov, T))   # name matched loosely
    assert rep.column == "D" and not rep.created_tab
    prior = gn.prior_month("Kalaimagal", "2026-12")
    assert prior.month == "2026-11"
    assert prior.savings_to_date == 11400 * 2
    assert prior.cash_in_hand == nov.get("weekly.cash_in_hand.W4")


def test_refuses_silent_overwrite_and_formula_cells(wbpath):
    gn = GNWorkbook(wbpath, T)
    vals = to_monthly(make_month(T), T)
    gn.write_month("Kalaimagal", "2026-10", vals, gn_name="PTK East", create_tab=True)
    changed = dict(vals, **{"Savings (Rs.)": 9999})
    with pytest.raises(WorkbookError, match="correction"):
        gn.write_month("Kalaimagal", "2026-10", changed)
    rep = gn.write_month("Kalaimagal", "2026-10", changed, overwrite=True)
    assert any(w.old == 11400 and w.new == 9999 for w in rep.writes)
    with pytest.raises(WorkbookError, match="formula"):
        gn.write_month("Kalaimagal", "2026-11", {"Total Income (Rs.)": 5})


def test_unknown_shg_suggests_close_names(wbpath):
    gn = GNWorkbook(wbpath, T)
    gn.write_month("Kalaimagal", "2026-10", to_monthly(make_month(T), T), gn_name="PTK East", create_tab=True)
    with pytest.raises(WorkbookError, match="Did you mean: Kalaimagal"):
        gn.write_month("Kalaimagel", "2026-11", {})
