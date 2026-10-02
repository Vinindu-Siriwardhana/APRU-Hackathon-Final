"""Round 1 API tests (owner: api): every endpoint and status guard, the Kalaimagal demo
path through the simulator, and the WhatsApp webhook (signature, de-duplication, odd
payloads, unsupported messages)."""
import hashlib
import hmac
import io
import json
import os
import tempfile
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

from round1_api_helpers import ROOT, seed

# app.api builds a default app at import time: point it at a throw-away folder first.
_DEFAULT = Path(tempfile.mkdtemp()) / "data"
_DEFAULT.mkdir()
seed(_DEFAULT / "gn_workbook.xlsx")
os.environ["SHG_DATA_DIR"] = str(_DEFAULT)
os.environ.pop("ANTHROPIC_API_KEY", None)

from app import api, whatsapp  # noqa: E402
from app import messages as M  # noqa: E402
from app.locking import replace  # noqa: E402

SYN = ROOT / "samples" / "synthetic"


@pytest.fixture
def client(tmp_path, monkeypatch):
    for k in ("WHATSAPP_APP_SECRET", "WHATSAPP_VERIFY_TOKEN", "WHATSAPP_TOKEN", "WHATSAPP_PHONE_ID",
              "SHG_DISABLE_SIMULATOR"):
        monkeypatch.delenv(k, raising=False)

    def seed_workbook(path):                     # demo reset re-seeds with the test helper
        tmp = path.with_name("seed.tmp.xlsx")
        seed(tmp)
        replace(tmp, path)
    monkeypatch.setattr(api, "seed_workbook", seed_workbook)
    wb = seed(tmp_path / "gn_workbook.xlsx")
    app = api.create_app(tmp_path / "data", wb, offline=True)
    c = TestClient(app)
    c.wb = wb
    c.data = tmp_path / "data"
    return c


def send(c, sender, name=None, lang=None, text=None, data=None):
    form = {"sender": sender}
    if lang:
        form["lang"] = lang
    if text is not None:
        form["text"] = text
        return c.post("/api/sim/message", data=form)
    payload = data if data is not None else (SYN / name).read_bytes()
    return c.post("/api/sim/message", data=form, files={"image": (name or "photo.jpg", payload, "image/jpeg")})


def kalaimagal(c, sender="kal", lang="ta"):
    send(c, sender, "kalaimagal_2026-10_p1.jpg", lang)
    r = send(c, sender, "kalaimagal_2026-10_p2.jpg").json()
    return r["submission_id"], r


# ================================================================ the demo path
def test_kalaimagal_demo_path(client):
    c = client
    sid, r = kalaimagal(c)
    assert r["status"] == "needs_review" and r["reply"] == M.t("checking", "ta") and r["reply_en"]
    s = c.get(f"/api/submissions/{sid}").json()
    top = s["validation"]["likely"][0]
    assert top["field"] == "weekly.interest_repaid.W3" and top["suggest"] == 310 and top["balances"] == 3
    assert s["short_labels"]["weekly.interest_repaid.W3"] == "Interest repaid, W3"
    assert s["monthly"]["Interest repayment (Rs.)"] == 1660                 # as it would be written now
    row = next(x for x in c.get("/api/submissions").json() if x["id"] == sid)
    assert row["likely_field"] == "weekly.interest_repaid.W3" and row["errors"] > 0

    s = c.post(f"/api/submissions/{sid}/fields/weekly.interest_repaid.W3", json={"value": "310"}).json()
    assert s["validation"]["likely"][0]["field"] == "weekly.savings.W2"
    s = c.post(f"/api/submissions/{sid}/fields/weekly.savings.W2", json={"value": 2800}).json()
    assert s["validation"]["likely"] == [] and s["actions"]["approve"]
    r = c.post(f"/api/submissions/{sid}/approve", json={"officer": "anu"}).json()
    assert r["submission"]["status"] == "awaiting_member" and r["delivered"]
    assert "அக்டோபர் 2026" in r["reply"] and "October 2026" in r["reply_en"]

    conv = c.get("/api/sim/conversation", params={"sender": "sim:kal"}).json()
    assert conv["lang"] == "ta" and conv["latest"] == {"id": sid, "status": "awaiting_member"}
    summary = conv["messages"][-1]
    assert summary["from"] == "bot" and summary["sid"] == sid
    assert summary["text"].startswith(M.t("summary_intro", "ta")) and summary["text_en"].startswith("Here is what we read")
    assert conv["messages"][0]["image"] == "kalaimagal_2026-10_p1.jpg"
    assert M.t("privacy", "ta") in conv["messages"][1]["text"]          # privacy note under the first reply
    assert all(M.t("privacy", "ta") not in m["text"] for m in conv["messages"][2:] if m["from"] == "bot")

    r = send(c, "kal", text="OK").json()
    assert r["status"] == "written" and r["reply"].startswith("✅") and "ரூ. 11,400" in r["reply"]
    assert r["reply_en"].startswith("✅ Your October 2026 report for Kalaimagal")
    s = c.get(f"/api/submissions/{sid}").json()
    assert s["write_report"]["sheet"] == "Kalaimagal" and s["write_report"]["column"] == "G"
    assert len(s["write_report"]["writes"]) == 16
    assert not [f for f, v in s["record"]["fields"].items() if v["status"] == "needs_review"]

    wb = openpyxl.load_workbook(io.BytesIO(c.get("/api/workbook").content))
    ws = wb["Kalaimagal"]
    approved = {w["cell"].split("!")[1]: w["new"] for w in s["write_report"]["writes"]}
    assert all(ws[cell].value == v for cell, v in approved.items())
    assert ws["G6"].value == 11400 and ws["G8"].value == 1160                 # savings, interest (310 not 810)


# ================================================================ endpoints + guards
def test_status_template_tabs_reasons(client):
    c = client
    st = c.get("/api/status").json()
    assert st["mode"] == "offline-demo" and st["simulator"] and not st["whatsapp"]
    assert len(c.get("/api/template").json()["summary_items"]) == 9
    assert c.get("/api/shg-tabs").json() == ["Kalaimagal", "Vasantham", "Sisila"]
    assert {r["code"] for r in c.get("/api/reject-reasons").json()} >= {"unclear", "missing_page", "other"}
    assert c.get("/api/groups").status_code == 200


def test_guards_return_409_with_a_human_reason(client):
    c = client
    r = send(c, "g", "kalaimagal_2026-10_p1.jpg").json()
    sid = r["submission_id"]
    for path, body in [("fields/weekly.savings.W1", {"value": 1}), ("approve", {}),
                       ("overrides", {"label": "Savings (Rs.)", "accept": True}), ("new-group", {}),
                       ("retry", {})]:
        res = c.post(f"/api/submissions/{sid}/{path}", json=body)
        assert res.status_code == 409, (path, res.text)
        assert res.json()["detail"].startswith("Can't")
    sid, _ = kalaimagal(c, "g")
    assert c.post(f"/api/submissions/{sid}/approve", json={}).status_code == 409          # checks failing
    assert c.post(f"/api/submissions/{sid}/approve", json={"force": True}).status_code == 400  # no reason
    assert c.post(f"/api/submissions/{sid}/overrides", json={"label": "Savings (Rs.)", "accept": False}).status_code == 404
    assert c.post(f"/api/submissions/{sid}/fields/weekly.bogus.W1", json={"value": 1}).status_code == 404
    for bad in (True, -5, 1e12, "abc", [1]):
        assert c.post(f"/api/submissions/{sid}/fields/weekly.savings.W1", json={"value": bad}).status_code == 400
    assert c.post(f"/api/submissions/{sid}/new-group", json={}).status_code == 409   # Kalaimagal has a tab
    r = c.post(f"/api/submissions/{sid}/approve", json={"force": True, "reason": "phoned the group"}).json()
    assert r["submission"]["status"] == "awaiting_member" and r["submission"]["forced"]["reason"] == "phoned the group"
    res = c.post(f"/api/submissions/{sid}/approve", json={})
    assert res.status_code == 409 and "waiting for her reply" in res.json()["detail"]
    send(c, "g", text="OK")
    for path, body in [("fields/weekly.savings.W1", {"value": 3100}), ("approve", {}), ("reject", {})]:
        assert c.post(f"/api/submissions/{sid}/{path}", json=body).status_code == 409
    assert c.get("/api/submissions/zzzzzzzzzz").status_code == 404
    assert c.get("/api/submissions/..%2f..%2fetc").status_code == 404


def test_member_correction_endpoints(client):
    c = client
    send(c, "m", "vasantham_2026-10_p1.jpg", "en")
    sid = send(c, "m", "vasantham_2026-10_p2.jpg").json()["submission_id"]
    assert c.get(f"/api/submissions/{sid}").json()["status"] == "awaiting_member"
    r = send(c, "m", text="5 12,000/-").json()
    assert r["status"] == "needs_review" and "noted" in r["reply"]
    s = c.get(f"/api/submissions/{sid}").json()
    corr = s["corrections"][0]
    assert corr["item"] == 5 and corr["value"] == 12000 and not corr["can_accept"]
    assert s["monthly"]["Savings (Rs.)"] != 12000                   # nothing changed behind her back
    res = c.post(f"/api/submissions/{sid}/overrides", json={"label": "Savings (Rs.)", "accept": True})
    assert res.status_code == 409 and "weekly cells" in res.json()["detail"]
    s = c.post(f"/api/submissions/{sid}/overrides", json={"label": "Savings (Rs.)", "accept": False}).json()
    assert s["corrections"] == [] and s["validation"]["likely"] == []
    assert c.post(f"/api/submissions/{sid}/approve", json={}).json()["submission"]["status"] == "awaiting_member"


def test_reject_retry_process_now(client):
    c = client
    sid, _ = kalaimagal(c, "rj", "si")
    r = c.post(f"/api/submissions/{sid}/reject", json={"reason": "unclear", "note": "too dark"}).json()
    assert r["submission"]["status"] == "rejected" and M.REJECT_REASONS["unclear"]["si"] in r["reply"]
    assert r["reply_en"] and "too dark" not in r["reply"]
    conv = c.get("/api/sim/conversation", params={"sender": "rj"}).json()
    assert conv["messages"][-1]["text"] == r["reply"]

    # a photo the offline demo doesn't know: failed, never "read"
    p1 = (SYN / "kalaimagal_2026-10_p1.jpg").read_bytes()
    import cv2, numpy as np
    resaved = cv2.imencode(".jpg", cv2.imdecode(np.frombuffer(p1, np.uint8), cv2.IMREAD_COLOR),
                           [cv2.IMWRITE_JPEG_QUALITY, 85])[1].tobytes()
    send(c, "u", "photo.jpg", data=resaved)
    r = send(c, "u", "kalaimagal_2026-10_p2.jpg").json()
    assert r["status"] == "failed" and r["reply"] == M.t("demo_unknown_photo", "en")
    row = next(x for x in c.get("/api/submissions").json() if x["id"] == r["submission_id"])
    assert row["failure"] and "sample photos" in row["failure"]
    rr = c.post(f"/api/submissions/{r['submission_id']}/retry", json={}).json()
    assert rr["submission"]["status"] == "failed" and rr["reply"] is None      # no second message to her

    r = send(c, "pn", "kalaimagal_2026-10_p1.jpg").json()
    pr = c.post(f"/api/submissions/{r['submission_id']}/process-now", json={"officer": "anu"}).json()
    assert pr["submission"]["status"] in ("needs_review", "awaiting_member")
    assert any(e.get("event") == "processed_without_page2" for e in pr["submission"]["log"])


def test_superseded_report_is_listed(client):
    c = client
    a, _ = kalaimagal(c, "a")
    b, _ = kalaimagal(c, "b")
    rows = {x["id"]: x for x in c.get("/api/submissions").json()}
    assert rows[a]["status"] == "superseded" and rows[a]["superseded_by"] == b


def test_member_text_states_in_her_language(client):
    c = client
    r = send(c, "t", text="hello").json()
    assert r["reply"].startswith(M.t("no_report", "en")) and r["submission_id"] is None
    kalaimagal(c, "t2", "si")
    r = send(c, "t2", text="OK").json()                               # officer still checking
    assert r["reply"] == M.t("officer_checking", "si") and r["reply_en"] == M.t("officer_checking", "en")


def test_uploads_are_checked(client):
    c = client
    assert send(c, "x", "a.jpg", data=b"").status_code == 400
    r = send(c, "x", "a.heic", data=b"not an image at all").json()
    assert r["reply"].startswith(M.t("bad_image", "en"))
    from PIL import Image
    buf = io.BytesIO()
    Image.new("1", (9000, 9000)).save(buf, "PNG")                     # 81 MP, tiny file
    r = send(c, "x", "bomb.png", data=buf.getvalue()).json()
    assert r["reply"].startswith(M.t("bad_image", "en"))
    assert send(c, "x", text="hi", lang="fr").status_code == 400
    assert send(c, "../../x", text="hi").status_code == 400


def test_paths_cannot_escape(client):
    c = client
    for url in ["/..%2f..%2fbackend%2fapp%2fwhatsapp.py", "//etc/passwd", "/api/sim/samples/..%5C..%5Cx.jpg",
                "/api/sim/samples/..%2fmanifest.json"]:
        r = c.get(url)
        assert "import" not in r.text and "root:" not in r.text, url
    assert c.get("/api/sim/samples/..%5C..%5Cx.jpg").status_code == 404
    sid, _ = kalaimagal(c, "p")
    assert c.get(f"/api/submissions/{sid}/crops/weekly.savings.W2").status_code == 200
    assert c.get(f"/api/submissions/{sid}/crops/..%5Csubmission").status_code == 404
    assert c.get(f"/api/submissions/{sid}/pages/1").status_code == 200
    assert c.get(f"/api/submissions/{sid}/pages/7").status_code == 404
    assert not (c.data / "submissions" / "abcdef0123").exists()
    c.get("/api/submissions/abcdef0123/crops/weekly.savings.W2")
    assert not (c.data / "submissions" / "abcdef0123").exists()       # a GET never creates folders


def test_simulator_can_be_disabled(client, monkeypatch):
    monkeypatch.setenv("SHG_DISABLE_SIMULATOR", "1")
    assert send(client, "x", text="hi").status_code == 404
    assert client.get("/api/sim/samples").status_code == 404


def test_reset_reseeds_then_clears(client, monkeypatch):
    c = client
    kalaimagal(c, "r")
    assert c.get("/api/submissions").json()
    monkeypatch.setattr(api, "seed_workbook", lambda path: (_ for _ in ()).throw(PermissionError("locked")))
    res = c.post("/api/demo/reset")
    assert res.status_code == 409 and "Excel" in res.json()["detail"]
    assert c.get("/api/submissions").json()                            # nothing cleared on failure
    monkeypatch.undo()
    monkeypatch.setattr(api, "seed_workbook", lambda path: (seed(path.with_name("s.xlsx")),
                                                            replace(path.with_name("s.xlsx"), path)))
    (c.data / "languages.json").write_text("{}", encoding="utf-8")
    assert c.post("/api/demo/reset").json() == {"ok": True}
    assert c.get("/api/submissions").json() == []
    assert c.get("/api/sim/conversation", params={"sender": "r"}).json()["messages"] == []
    assert not (c.data / "languages.json").exists()


# ================================================================ WhatsApp webhook
def wa_payload(*msgs):
    return {"object": "whatsapp_business_account",
            "entry": [{"changes": [{"value": {"messages": list(msgs)}}]}]}


def sign(body: bytes, secret="s3cret"):
    return "sha256=" + hmac.new(secret.encode(), body, hashlib.sha256).hexdigest()


@pytest.fixture
def wa(client, monkeypatch):
    sent = []
    monkeypatch.setenv("WHATSAPP_APP_SECRET", "s3cret")
    monkeypatch.setenv("WHATSAPP_TOKEN", "t")
    monkeypatch.setenv("WHATSAPP_PHONE_ID", "p")
    monkeypatch.setattr(whatsapp, "send_text", lambda to, text: sent.append((to, text)) or {})
    monkeypatch.setattr(whatsapp, "download_media", lambda mid: (SYN / mid).read_bytes())

    def post(*msgs, signature=True, raw=None):
        body = raw if raw is not None else json.dumps(wa_payload(*msgs)).encode()
        headers = {"Content-Type": "application/json"}
        if signature:
            headers["X-Hub-Signature-256"] = sign(body) if signature is True else signature
        return client.post("/webhook/whatsapp", content=body, headers=headers)
    client.post_wa, client.sent = post, sent
    return client


def test_webhook_verify_handshake(client, monkeypatch):
    q = {"hub.mode": "subscribe", "hub.challenge": "12345"}
    assert client.get("/webhook/whatsapp", params=q).status_code == 403            # no token configured
    monkeypatch.setenv("WHATSAPP_VERIFY_TOKEN", "tok")
    assert client.get("/webhook/whatsapp", params={**q, "hub.verify_token": "nope"}).status_code == 403
    r = client.get("/webhook/whatsapp", params={**q, "hub.verify_token": "tok"})
    assert r.status_code == 200 and r.text == "12345"


def test_webhook_signature(wa):
    msg = {"from": "94771234567", "id": "wamid.1", "type": "text", "text": {"body": "hello"}}
    assert wa.post_wa(msg, signature=False).status_code == 401
    assert wa.post_wa(msg, signature="sha256=" + "0" * 64).status_code == 401
    assert wa.sent == []
    r = wa.post_wa(msg)
    assert r.status_code == 200 and r.json()["received"] == 1
    assert wa.sent and wa.sent[0][0] == "94771234567" and M.t("no_report", "en") in wa.sent[0][1]


def test_webhook_deduplicates_redelivery(wa):
    p1 = {"from": "9477", "id": "wamid.p1", "type": "image", "image": {"id": "vasantham_2026-10_p1.jpg"}}
    p2 = {"from": "9477", "id": "wamid.p2", "type": "image", "image": {"id": "vasantham_2026-10_p2.jpg"}}
    assert wa.post_wa(p1).json()["duplicates"] == 0
    wa.post_wa(p2)
    assert wa.post_wa(p2).json()["duplicates"] == 1                     # Meta retries page 2
    subs = wa.get("/api/submissions").json()
    assert len(subs) == 1 and subs[0]["status"] == "awaiting_member"
    ok = {"from": "9477", "id": "wamid.ok", "type": "text", "text": {"body": "OK"}}
    wa.post_wa(ok)
    wa.post_wa(ok)
    assert wa.get("/api/submissions").json()[0]["status"] == "written"
    assert sum("✅" in t for _, t in wa.sent) == 1                      # one receipt


def test_webhook_odd_payloads_and_unsupported_types(wa):
    assert wa.post_wa(raw=b"not json").status_code == 400
    for payload in [{}, {"entry": ["x"]}, {"entry": [{"changes": [{"value": {"messages": [{"type": "text"}]}}]}]},
                    {"entry": [{"changes": [{"value": {"statuses": [{"status": "read"}]}}]}]}]:
        body = json.dumps(payload).encode()
        assert wa.post_wa(raw=body).status_code == 200
    r = wa.post_wa({"from": "9478", "id": "wamid.a", "type": "audio", "audio": {"id": "x"}},
                   {"from": "9478", "id": "wamid.t", "type": "text", "text": {}})
    assert r.status_code == 200
    assert [t for _, t in wa.sent if t.startswith(M.t("unsupported", "en"))]
    pdf = {"from": "9479", "id": "wamid.d", "type": "document", "document": {"id": "x", "mime_type": "application/pdf"}}
    wa.post_wa(pdf)
    assert wa.sent[-1][0] == "9479" and wa.sent[-1][1].startswith(M.t("unsupported", "en"))
    stop = {"from": "9479", "id": "wamid.s", "type": "text", "text": {"body": "STOP"}}
    wa.post_wa(stop)
    assert wa.sent[-1][1] == M.t("stopped", "en")
