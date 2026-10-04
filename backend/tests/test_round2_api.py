"""Round 2 API/pipeline tests (owner: api).

- reading races: an officer reject or a member STOP while the report is being read wins
- the quality check runs outside the global lock; trivial GETs are async
- member replies go to the summary that is waiting for them (two open months)
- delivery status on every bot message, the undelivered count, Resend, Meta status callbacks
- WhatsApp webhook edge cases (non-ASCII signature, deep JSON, unsigned posts in live mode,
  ids remembered only after processing)
- a new group's opening totals: started this month vs the mother book, opening_inconsistent,
  the workbook's column C and Palmera's own cash check after a LibreOffice recalculation
- "Send anyway" and "checked as written" limits, month outside the tab, 500 -> 400
- summary_items (what she confirms), officer-friendly failure reasons and 4xx details
- the Kalaimagal demo path end to end
"""
import hashlib
import hmac
import inspect
import io
import json
import shutil
import subprocess
import threading
import time
from pathlib import Path

import httpx
import openpyxl
import pytest

from test_round1_api import SYN, client, kalaimagal, seed, send, sign, wa_payload  # noqa: F401  (fixture)

from app import api, whatsapp
from app import messages as M
from app import pipeline as P
from app.api import decode_image
from app.conversations import SeenMessages
from app.demo_reader import DemoReader, SampleIndex
from app.locking import lock
from app.pipeline import Pipeline, StateError, Status, Store
from app.schema import load_template
from app.validation import OPENING_LABELS

T = load_template()
SAV, OUT = "Total savings to date (Rs.)", "Total Loans outstanding to date (Rs.)"


def img(name):
    return decode_image((SYN / name).read_bytes())


def view(c, sid):
    return c.get(f"/api/submissions/{sid}").json()


def errors(s):
    return {i["rule"] for i in s["validation"]["issues"] if i["severity"] == "error"}


def vasantham(c, sender="v", lang="en"):
    send(c, sender, "vasantham_2026-10_p1.jpg", lang)
    r = send(c, sender, "vasantham_2026-10_p2.jpg").json()
    return r["submission_id"], r


# ================================================================ reading races (F1 / P1-4)
class Gate:
    """A reader that blocks inside read() until the test lets it go."""

    def __init__(self, inner):
        self.inner, self.started, self.go = inner, threading.Event(), threading.Event()

    def for_submission(self, sub):
        r, gate = self.inner.for_submission(sub), self

        class Blocked:
            def read(self, images, prompt, schema):
                gate.started.set()
                assert gate.go.wait(20)
                return r.read(images, prompt, schema)
        return Blocked()


@pytest.fixture
def gated(tmp_path):
    gate = Gate(DemoReader(T, SampleIndex(SYN)))
    p = Pipeline(Store(tmp_path / "data"), seed(tmp_path / "gn.xlsx"), gate, T)
    return p, gate


def _read_in_background(p, gate, sender):
    p.on_image(sender, img("vasantham_2026-10_p1.jpg"), lang="en", source="vasantham_2026-10_p1.jpg")
    out = {}
    t = threading.Thread(target=lambda: out.update(r=p.on_image(sender, img("vasantham_2026-10_p2.jpg"),
                                                               source="vasantham_2026-10_p2.jpg")))
    t.start()
    assert gate.started.wait(20)
    sid = next(s.id for s in p.store.for_sender(sender))
    assert sid in p._busy
    return t, out, sid


def _nothing_sent_after(p, sender, sub):
    texts = [m["text"] for m in p.conversations.load(sender).messages if m["from"] == "bot"]
    assert not any(M.t("summary_intro", "en") in t or M.t("checking", "en") in t for t in texts)
    assert not any(M.t("summary_intro", "en") in e.get("text", "") for e in sub.log)
    assert any(e.get("event") == "reading_discarded" for e in sub.log)


def test_officer_reject_during_reading_wins(gated):
    p, gate = gated
    t, out, sid = _read_in_background(p, gate, "rd")
    p.officer_reject(p.store.load(sid), "anu", "unclear")
    gate.go.set()
    t.join(20)
    sub, reply = out["r"]
    assert reply is None                                   # she was already told; nothing more
    sub = p.store.load(sid)
    assert sub.status == Status.rejected and sub.rejected_reason == "unclear" and sub.record is None
    _nothing_sent_after(p, "rd", sub)
    _, reply = p.member_reply("rd", "OK")                  # her OK can't write the rejected report
    assert reply == M.t("no_report", "en")
    ws = openpyxl.load_workbook(p.workbook_path)["Vasantham"]
    assert ws["G6"].value is None


def test_member_stop_during_reading_wins(gated):
    p, gate = gated
    t, out, sid = _read_in_background(p, gate, "st")
    _, stop = p.on_text("st", "STOP")
    assert stop == M.t("stopped", "en")
    gate.go.set()
    t.join(20)
    assert out["r"][1] is None
    sub = p.store.load(sid)
    assert sub.status == Status.rejected and sub.rejected_reason == "member_opted_out"
    _nothing_sent_after(p, "st", sub)
    assert not any(s.status in P.OPEN for s in p.store.all())   # not back in the officer's queue


def test_retry_rejected_during_reading_is_discarded(gated, tmp_path):
    p, gate = gated
    gate.go.set()
    p.on_image("rt", img("kalaimagal_2026-10_p1.jpg"), source=None)       # unknown photo -> failed
    sub, _ = p.on_image("rt", img("kalaimagal_2026-10_p2.jpg"), source="kalaimagal_2026-10_p2.jpg")
    assert sub.status == Status.failed
    gate.go.clear(); gate.started.clear()
    p.store.load(sub.id)
    s = p.store.load(sub.id); s.page_sources[1] = "kalaimagal_2026-10_p1.jpg"; p.store.save(s)
    out = {}
    t = threading.Thread(target=lambda: out.update(r=p.retry(sub.id, "anu")))
    t.start()
    assert gate.started.wait(20)
    p.officer_reject(p.store.load(sub.id), "anu", "unclear")
    gate.go.set(); t.join(20)
    assert out["r"] is None and p.store.load(sub.id).status == Status.rejected


# ================================================================ lock + async GETs (F2)
def test_quality_check_runs_outside_the_lock(gated, monkeypatch):
    p, gate = gated
    gate.go.set()
    inside, release = threading.Event(), threading.Event()
    real = P.check

    def slow_check(image, loc):
        inside.set()
        assert release.wait(20)
        return real(image, loc)
    monkeypatch.setattr(P, "check", slow_check)
    t = threading.Thread(target=p.receive_image, args=("q", img("vasantham_2026-10_p1.jpg")))
    t.start()
    assert inside.wait(20)
    got = lock.acquire(timeout=0.5)                       # the dashboard can still read
    assert got
    lock.release()
    release.set(); t.join(20)
    sub = p.store.for_sender("q")[0]
    assert sub.pages == {1: "page1.jpg"} and (p.store.dir(sub.id) / "page1.jpg").stat().st_size > 10_000


def test_trivial_gets_are_async(client):
    eps = {r.path: r.endpoint for r in client.app.routes if hasattr(r, "endpoint")}
    for path in ("/api/status", "/api/template", "/api/reject-reasons"):
        assert inspect.iscoroutinefunction(eps[path]), path


# ================================================================ routing (P2-4)
def test_ok_goes_to_the_summary_waiting_for_her(client):
    c = client
    oct_id, r = vasantham(c, "two")
    assert r["status"] == "awaiting_member"
    nov = send(c, "two", "kalaimagal_2026-11_p1.jpg").json()                # next report started
    assert nov["status"] == "collecting" and nov["submission_id"] != oct_id
    r = send(c, "two", text="hello").json()                                # other text: newest report
    assert r["submission_id"] == nov["submission_id"] and r["reply"] == M.t("waiting_page2", "en")
    r = send(c, "two", text="OK").json()
    assert r["submission_id"] == oct_id and r["status"] == "written" and r["reply"].startswith("✅")
    assert view(c, nov["submission_id"])["status"] == "collecting"


def test_correction_goes_to_the_summary_too(client):
    c = client
    oct_id, _ = vasantham(c, "two-c")
    send(c, "two-c", "kalaimagal_2026-11_p1.jpg")
    r = send(c, "two-c", text="5 12000").json()
    assert r["submission_id"] == oct_id and r["status"] == "needs_review"


# ================================================================ 500 -> 400 (F6)
def test_huge_numbers_are_a_400_not_a_500(client):
    c = client
    sid, _ = kalaimagal(c, "big")
    for fid in ("weekly.savings.W1", "header.month_year", "header.shg_name", "weekly.attendance.W1",
                "header.total_members"):
        r = c.post(f"/api/submissions/{sid}/fields/{fid}", content=f'{{"value": {10 ** 400}}}',
                   headers={"Content-Type": "application/json"})
        assert r.status_code == 400, (fid, r.text)
        assert r.json()["detail"].endswith(".")
    r = c.post(f"/api/submissions/{sid}/fields/weekly.savings.W1", json={"value": "j"})
    assert r.status_code == 400 and r.json()["detail"] == "“j” isn't a number. Enter it in digits, like 1200."
    r = c.post(f"/api/submissions/{sid}/fields/header.month_year", json={"value": "0001-01"})
    assert r.status_code == 400 and "between 2000 and 2100" in r.json()["detail"]
    r = c.post(f"/api/submissions/{sid}/fields/weekly.meeting_held.W1", json={"value": "maybe"})
    assert r.status_code == 400 and r.json()["detail"] == "Choose Yes or No."


# ================================================================ 4xx details are sentences
def test_every_4xx_detail_is_a_sentence(client):
    c = client
    sid, _ = kalaimagal(c, "s4")
    calls = [c.get("/api/submissions/abcdef0123"), c.get("/api/nope"),
             c.get(f"/api/submissions/{sid}/pages/7"), c.get(f"/api/submissions/{sid}/crops/x.y"),
             c.post(f"/api/submissions/{sid}/fields/weekly.bogus.W1", json={"value": 1}),
             c.post(f"/api/submissions/{sid}/overrides", json={"label": "Savings (Rs.)", "accept": False}),
             c.post(f"/api/submissions/{sid}/overrides", json={"accept": "perhaps"}),
             c.post(f"/api/submissions/{sid}/approve", json={}), c.post(f"/api/submissions/{sid}/retry"),
             c.post(f"/api/submissions/{sid}/new-group", json={}),
             send(c, "s4", text="hi", lang="fr"), c.post("/api/sim/message", data={"sender": "s4"}), c.get("/webhook/whatsapp"),
             c.get("/api/sim/samples/nope.jpg"), c.post("/api/submissions")]
    for r in calls:
        assert 400 <= r.status_code < 500, r.text
        d = r.json()["detail"]
        assert isinstance(d, str) and d[0].isupper() and d.rstrip("”").endswith((".", "?")), (r.status_code, d)
        assert d not in ("Not Found", "Method Not Allowed") and "ANTHROPIC" not in d


# ================================================================ failure wording (UI P2-1)
def test_unknown_photo_failure_is_officer_friendly(client):
    import cv2
    import numpy as np
    c = client
    p1 = (SYN / "kalaimagal_2026-10_p1.jpg").read_bytes()
    resaved = cv2.imencode(".jpg", cv2.imdecode(np.frombuffer(p1, np.uint8), cv2.IMREAD_COLOR),
                           [cv2.IMWRITE_JPEG_QUALITY, 85])[1].tobytes()
    send(c, "uf", "photo.jpg", data=resaved)
    r = send(c, "uf", "kalaimagal_2026-10_p2.jpg").json()
    s = view(c, r["submission_id"])
    assert s["status"] == "failed" and s["failure_kind"] == "unknown_sample"
    assert s["failure"] == ("This photo isn't one of the demo samples, so offline demo mode can't read it. "
                            "Ask the member for a new photo.")
    assert not s["actions"]["retry"] and s["actions"]["reject"]           # "Ask for a new photo"
    row = next(x for x in c.get("/api/submissions").json() if x["id"] == s["id"])
    assert row["failure_kind"] == "unknown_sample"


# ================================================================ summary_items (UI P1-2, P1-5)
def test_summary_items_are_her_whatsapp_summary(client):
    c = client
    sid, r = vasantham(c, "si-items", "ta")
    s = view(c, sid)
    items = s["summary_items"]
    assert [i["n"] for i in items] == list(range(1, 11)) and len(items) <= 10
    assert items[4]["label_en"] == "Savings (Rs.)" and items[4]["value"] == 12400
    assert items[9]["monthly_label"] == "Number of SHG members" and items[9]["value"] == 19
    assert items[0]["monthly_label"] is None and items[1]["display_en"] == "October 2026"
    summary = r["reply"]
    for i in items:                                        # every numbered line is in what she got
        assert f"{i['n']}. {i['label']}: {i['display']}" in summary
    monthly = {row["label"] for row in c.get("/api/template").json()["monthly"]}
    assert {i["monthly_label"] for i in items if i["monthly_label"]} <= monthly


def test_member_count_can_be_corrected_but_not_to_nonsense(client):
    c = client
    sid, _ = vasantham(c, "mem")
    r = send(c, "mem", text="10 5000").json()
    assert r["reply"] == M.t("not_understood_value", "en", n=10)
    r = send(c, "mem", text="10 18").json()
    assert r["status"] == "needs_review"
    corr = view(c, sid)["corrections"][0]
    assert corr["item"] == 10 and corr["value"] == 18 and corr["can_accept"] and corr["cell"] == "header.total_members"


# ================================================================ Send anyway / checked-as-written limits
def test_send_anyway_cannot_skip_unknown_group_or_month_problems(client):
    c = client
    sid, _ = vasantham(c, "fa")
    c.post(f"/api/submissions/{sid}/fields/header.shg_name", json={"value": "Nilmini"})
    r = c.post(f"/api/submissions/{sid}/approve", json={"force": True, "reason": "phoned"})
    assert r.status_code == 409 and r.json()["detail"].startswith("Can't send anyway: this group has no tab")

    sid, _ = vasantham(c, "fb")
    s = c.post(f"/api/submissions/{sid}/fields/header.month_year", json={"value": "2028-10"}).json()
    assert "month_out_of_range" in errors(s)
    issue = next(i for i in s["validation"]["issues"] if i["rule"] == "month_out_of_range")
    assert issue["message"].startswith("The form says October 2028, but the Vasantham tab in the workbook only "
                                       "has columns for June 2026 to May 2028.")
    s = c.post(f"/api/submissions/{sid}/fields/header.month_year", json={"value": "2028-10"}).json()
    assert "month_out_of_range" in errors(s)                         # confirming it changes nothing
    r = c.post(f"/api/submissions/{sid}/approve", json={"force": True, "reason": "phoned"})
    assert r.status_code == 409 and "Can't send anyway" in r.json()["detail"]

    sid, _ = vasantham(c, "fc")
    s = c.post(f"/api/submissions/{sid}/fields/header.month_year", json={"value": "2026-09"}).json()
    assert "month_already_recorded" in errors(s)
    r = c.post(f"/api/submissions/{sid}/approve", json={"force": True, "reason": "phoned"})
    assert r.status_code == 409 and "confirm it as a correction" in r.json()["detail"]


def test_checked_as_written_never_waives_palmeras_maximum(client):
    c = client
    sid, _ = vasantham(c, "mx")
    c.post(f"/api/submissions/{sid}/fields/header.total_members", json={"value": 55})
    s = c.post(f"/api/submissions/{sid}/fields/header.total_members", json={"value": 55}).json()
    assert s["record"]["fields"]["header.total_members"]["status"] == "officer_confirmed"
    assert "above_max" in errors(s)
    assert c.post(f"/api/submissions/{sid}/approve", json={}).status_code == 409


# ================================================================ new group + opening totals (P1-2)
def _rename(c, sender):
    sid, _ = vasantham(c, sender)
    s = c.post(f"/api/submissions/{sid}/fields/header.shg_name", json={"value": "Nilmini"}).json()
    assert "unknown_shg" in errors(s) and s["actions"]["new_group"]
    return sid, s


def mother_book(suggest):
    """Opening totals consistent with the form: Vasantham's September (cash Rs 9,340)."""
    o = {k: 0 for k in OPENING_LABELS}
    o[SAV], o[OUT] = suggest[SAV], suggest[OUT]                      # 52,600 and 42,000
    o["Total Interest repayments to date (Rs.)"] = 8000
    o["Total Principal loan repayments to date (Rs.)"] = 30000
    o["Total Loans distributed to date (Rs.)"] = 30000 + suggest[OUT]
    o["Total Other expenses to date (Rs.)"] = 8000 + (suggest[SAV] - suggest[OUT]) - 9340
    return o


def test_new_group_that_did_not_start_this_month_is_caught(client):
    c = client
    sid, s = _rename(c, "ng1")
    assert s["opening_suggestion"][SAV] == 52600 and s["opening_suggestion"][OUT] == 42000
    s = c.post(f"/api/submissions/{sid}/new-group", json={"started_this_month": True}).json()
    assert s["started_this_month"] and s["opening"] == {k: 0 for k in OPENING_LABELS}
    issue = next(i for i in s["validation"]["issues"] if i["rule"] == "opening_inconsistent")
    assert issue["severity"] == "error" and issue["message"].startswith(
        "The form says total savings to date is Rs 65,000, but the group's only savings this month is Rs 12,400, "
        "so it did not start this month. Enter its totals from the mother book.")
    assert s["actions"]["opening"] and not s["actions"]["new_group"]
    r = c.post(f"/api/submissions/{sid}/approve", json={"force": True, "reason": "x"})
    assert r.status_code == 409 and "opening totals" in r.json()["detail"]
    # a mother-book figure that doesn't match the form is caught too
    o = mother_book(s["opening_suggestion"] or {SAV: 52600, OUT: 42000})
    s = c.post(f"/api/submissions/{sid}/new-group", json={"opening": {**o, SAV: 50000}}).json()
    issue = next(i for i in s["validation"]["issues"] if i["rule"] == "opening_inconsistent")
    assert "Rs 52,600 before October 2026, but Rs 50,000 was entered" in issue["message"]


def test_new_group_opening_body_is_checked(client):
    c = client
    sid, s = _rename(c, "ng2")
    o = mother_book(s["opening_suggestion"])
    for body, words in [({}, "started this month"), ({"opening": {SAV: 1}}, "Missing"),
                        ({"opening": {**o, "Bogus": 1}}, "isn't one of the nine"),
                        ({"opening": {**o, SAV: -5}}, "between 0"), ({"opening": {**o, SAV: "lots"}}, "isn't an amount"),
                        ({"started_this_month": True, "opening": o}, "not both")]:
        r = c.post(f"/api/submissions/{sid}/new-group", json=body)
        assert r.status_code == 400 and words in r.json()["detail"], (body, r.text)
    assert view(c, sid)["new_group"] is False


def _recalc(xlsx: Path, out: Path):
    out.mkdir(parents=True, exist_ok=True)
    subprocess.run(["soffice", f"-env:UserInstallation=file://{out / 'profile'}", "--headless", "--calc",
                    "--convert-to", "xlsx", "--outdir", str(out), str(xlsx)],
                   check=True, capture_output=True, timeout=180)
    return openpyxl.load_workbook(out / xlsx.name, data_only=True)


def test_new_group_with_mother_book_totals_end_to_end(client, tmp_path):
    c = client
    sid, s = _rename(c, "ng3")
    o = mother_book(s["opening_suggestion"])
    s = c.post(f"/api/submissions/{sid}/new-group", json={"opening": {k: f"{v:,}" for k, v in o.items()}}).json()
    assert s["opening"] == o and not errors(s), s["validation"]["issues"]
    assert "unknown_shg" in {i["rule"] for i in s["validation"]["issues"]}       # as a warning
    r = c.post(f"/api/submissions/{sid}/approve", json={}).json()
    assert r["submission"]["status"] == "awaiting_member"
    r = send(c, "ng3", text="OK").json()
    assert r["status"] == "written"
    assert "Your group's total savings so far: Rs. 65,000." in r["reply"]           # backed by the opening
    s = view(c, sid)
    assert s["write_report"]["created_tab"] and s["write_report"]["column"] == "C"

    content = c.get("/api/workbook").content
    wb = openpyxl.load_workbook(io.BytesIO(content))
    ws = wb["Nilmini"]
    flows = {"Total savings to date (Rs.)": 12400, OUT: 51700 - 42000}
    assert ws["C21"].value == o[SAV] + flows[SAV] == 65000
    assert ws["C29"].value == 51700
    assert ws["C19"].value == 13030
    if not shutil.which("soffice"):
        pytest.skip("LibreOffice not installed: Palmera's cash check not recalculated")
    x = tmp_path / "gn.xlsx"
    x.write_bytes(content)
    calc = _recalc(x, tmp_path / "calc")["Nilmini"]
    assert calc["C20"].value == pytest.approx(calc["C19"].value) == 13030            # Palmera's own check


def test_new_group_started_this_month_end_to_end(client):
    """A form whose to-date figures are this month's own: started_this_month is consistent."""
    c = client
    sid, _ = vasantham(c, "ng4")
    s = view(c, sid)
    fields = s["record"]["fields"]
    sav = out = 0
    edits = {"header.shg_name": "Nilmini"}
    for wk in ("W1", "W2", "W3", "W4"):
        g = lambda k: (fields.get(f"weekly.{k}.{wk}") or {}).get("value") or 0
        sav += g("savings") - 0
        out += g("loans_distributed") - g("principal_repaid") - g("write_offs")
        edits[f"weekly.savings_to_date.{wk}"] = sav
        edits[f"weekly.loans_outstanding.{wk}"] = out if out >= 0 else None
    for fid, v in edits.items():
        s = c.post(f"/api/submissions/{sid}/fields/{fid}", json={"value": v}).json()
    if any(v is None for v in edits.values()):
        pytest.skip("this sample repays more than it lends in week 1")
    s = c.post(f"/api/submissions/{sid}/new-group", json={"started_this_month": True}).json()
    assert "opening_inconsistent" not in errors(s)


# ================================================================ delivery status + resend (F3)
@pytest.fixture
def live(client, monkeypatch):
    """WhatsApp configured; send_text controllable: c.fail = None (works) or an exception."""
    monkeypatch.setenv("WHATSAPP_APP_SECRET", "s3cret")
    monkeypatch.setenv("WHATSAPP_TOKEN", "t")
    monkeypatch.setenv("WHATSAPP_PHONE_ID", "p")
    client.sent, client.fail, n = [], None, iter(range(10 ** 6))

    def send_text(to, text):
        if client.fail:
            raise client.fail
        client.sent.append((to, text))
        return {"messages": [{"id": f"wamid.out{next(n)}"}]}
    monkeypatch.setattr(whatsapp, "send_text", send_text)
    monkeypatch.setattr(whatsapp, "download_media", lambda mid: (SYN / mid).read_bytes())

    def post(*msgs, raw=None, statuses=None):
        if raw is None:
            payload = wa_payload(*msgs)
            if statuses:
                payload["entry"][0]["changes"][0]["value"]["statuses"] = statuses
            raw = json.dumps(payload).encode()
        return client.post("/webhook/whatsapp", content=raw,
                           headers={"Content-Type": "application/json", "X-Hub-Signature-256": sign(raw)})
    client.post_wa = post
    return client


def window_error():
    req = httpx.Request("POST", "https://graph.facebook.com/v21.0/p/messages")
    resp = httpx.Response(400, json={"error": {"code": 131047, "message": "Re-engagement message"}}, request=req)
    return httpx.HTTPStatusError("400", request=req, response=resp)


def photo_msg(i, name, sender="9477"):
    return {"from": sender, "id": f"wamid.in{i}", "type": "image", "image": {"id": name}}


def test_failed_sends_are_recorded_and_resent_in_order(live):
    c = live
    c.fail = window_error()
    c.post_wa(photo_msg(1, "vasantham_2026-10_p1.jpg"))
    c.post_wa(photo_msg(2, "vasantham_2026-10_p2.jpg"))
    row = c.get("/api/submissions").json()[0]
    assert row["status"] == "awaiting_member" and row["undelivered"] == 2           # "send page 2", summary
    s = view(c, row["id"])
    bad = [e for e in s["log"] if e.get("from") == "bot"]
    assert all(e["delivered"] is False for e in bad) and s["actions"]["resend"] and s["undelivered"] == 2
    assert bad[0]["delivery_error"].startswith("Not delivered: more than 24 hours have passed")
    conv = c.app.state.pipeline.conversations.load("9477")
    assert [m["delivered"] for m in conv.messages if m["from"] == "bot"] == [False, False]

    r = c.post(f"/api/submissions/{row['id']}/resend", json={"officer": "anu"}).json()   # still failing
    assert r["resent"] == 0 and r["failed"] == 2 and "24 hours" in r["delivery_error"]
    c.fail = None
    r = c.post(f"/api/submissions/{row['id']}/resend", json={"officer": "anu"}).json()
    assert r["resent"] == 2 and r["failed"] == 0 and r["submission"]["undelivered"] == 0
    assert [t.split("\n")[0][:30] for _, t in c.sent] == [M.t("ask_page2", "en")[:30], M.t("summary_intro", "en")[:30]]
    assert M.t("privacy", "en") in c.sent[0][1]                    # sent as it was meant to go out
    assert c.get("/api/submissions").json()[0]["undelivered"] == 0
    r = c.post(f"/api/submissions/{row['id']}/resend", json={}).json()
    assert (r["resent"], r["failed"]) == (0, 0) and len(c.sent) == 2              # nothing sent twice


def test_officer_action_records_delivery(live):
    c = live
    c.post_wa(photo_msg(3, "kalaimagal_2026-10_p1.jpg", "9480"))
    c.post_wa(photo_msg(4, "kalaimagal_2026-10_p2.jpg", "9480"))
    sid = c.get("/api/submissions").json()[0]["id"]
    c.fail = RuntimeError("connection reset")
    r = c.post(f"/api/submissions/{sid}/reject", json={"reason": "unclear"}).json()
    assert r["delivered"] is False and "could not be reached" in r["delivery_error"]
    entry = [e for e in r["submission"]["log"] if e.get("from") == "bot"][-1]
    assert entry["delivered"] is False and r["submission"]["undelivered"] == 1


def test_meta_failed_status_marks_the_message(live):
    c = live
    c.post_wa(photo_msg(5, "vasantham_2026-10_p1.jpg", "9481"))
    conv = c.app.state.pipeline.conversations.load("9481")
    m = [m for m in conv.messages if m["from"] == "bot"][-1]
    assert m["delivered"] is True and m["wamid"].startswith("wamid.out")
    st = [{"id": m["wamid"], "recipient_id": "9481", "status": "failed",
           "errors": [{"code": 131047, "title": "Re-engagement message"}]}]
    assert c.post_wa(statuses=st).status_code == 200
    conv = c.app.state.pipeline.conversations.load("9481")
    m2 = conv.bot_message(mid=m["mid"])
    assert m2["delivered"] is False and "24 hours" in m2["delivery_error"]
    assert c.get("/api/submissions").json()[0]["undelivered"] == 1


def test_simulator_messages_have_null_delivery(client):
    c = client
    sid, _ = kalaimagal(c, "simd")
    s = view(c, sid)
    assert all(e["delivered"] is None for e in s["log"] if e.get("from") == "bot")
    assert s["undelivered"] == 0 and not s["actions"]["resend"]
    msgs = c.get("/api/sim/conversation", params={"sender": "simd"}).json()["messages"]
    assert all("mid" in m and m["delivered"] is None for m in msgs if m["from"] == "bot")


def test_describe_send_error():
    assert whatsapp.describe_send_error(window_error()).startswith("Not delivered: more than 24 hours")
    assert "token" in whatsapp.describe_error(190)
    assert whatsapp.describe_send_error(RuntimeError("WHATSAPP_TOKEN is not set")).startswith("Not sent")


# ================================================================ webhook edge cases (F7, F8, F11)
def test_webhook_non_ascii_signature_and_deep_json(live):
    c = live
    body = json.dumps(wa_payload()).encode()
    r = c.post("/webhook/whatsapp", content=body,
               headers={"Content-Type": "application/json", "X-Hub-Signature-256": b"sha256=\xe9\xe9"})
    assert r.status_code == 401
    assert not whatsapp.signature_ok(body, "sha256=é" + "0" * 63)
    deep = b"[" * 5000 + b"]" * 5000
    r = c.post_wa(raw=deep)
    assert r.status_code == 400 and r.json()["detail"].endswith(".")


def test_unsigned_posts_are_refused_in_live_mode(client, monkeypatch):
    c = client
    monkeypatch.delenv("WHATSAPP_APP_SECRET", raising=False)
    monkeypatch.setenv("WHATSAPP_TOKEN", "t")
    monkeypatch.setenv("WHATSAPP_PHONE_ID", "p")
    monkeypatch.setattr(whatsapp, "send_text", lambda to, text: {})
    body = json.dumps(wa_payload({"from": "9490", "id": "wamid.u", "type": "text", "text": {"body": "OK"}})).encode()
    r = c.post("/webhook/whatsapp", content=body, headers={"Content-Type": "application/json"})
    assert r.status_code == 503 and "WHATSAPP_APP_SECRET" in r.json()["detail"]
    st = c.get("/api/status").json()
    assert st["webhook"] == "refused" and "WHATSAPP_APP_SECRET" in st["webhook_warning"]
    monkeypatch.setenv("WHATSAPP_ALLOW_UNSIGNED", "1")
    assert c.post("/webhook/whatsapp", content=body, headers={"Content-Type": "application/json"}).status_code == 200
    assert c.get("/api/status").json()["webhook"] == "unsigned"
    monkeypatch.setenv("WHATSAPP_APP_SECRET", "s3cret")
    assert c.get("/api/status").json()["webhook"] == "signed"


def test_seen_ids_are_saved_only_after_processing(tmp_path):
    s = SeenMessages(tmp_path)
    assert s.claim("a") and not s.claim("a")                 # in flight: a re-delivery is ignored
    assert SeenMessages(tmp_path).claim("a")                 # ...but nothing is on disk yet
    s.done("a")
    assert not SeenMessages(tmp_path).claim("a") and not s.claim("a")
    assert s.claim("b"); s.release("b"); assert s.claim("b")
    assert s.claim(None) and s.claim(None)


def test_webhook_remembers_processed_ids(live):
    c = live
    c.post_wa({"from": "9491", "id": "wamid.x1", "type": "text", "text": {"body": "hi"}})
    assert "wamid.x1" in json.loads((c.data / "whatsapp_seen.json").read_text())
    assert c.post_wa({"from": "9491", "id": "wamid.x1", "type": "text", "text": {"body": "hi"}}).json()["duplicates"] == 1


def test_pages_and_crops_after_reset_are_404(client):
    c = client
    sid, _ = kalaimagal(c, "gone")
    shutil.rmtree(c.data / "submissions" / sid / "crops")
    (c.data / "submissions" / sid / "page1.jpg").unlink()
    for url in (f"/api/submissions/{sid}/pages/1", f"/api/submissions/{sid}/crops/weekly.savings.W2"):
        r = c.get(url)
        assert r.status_code == 404 and "Restart demo" in r.json()["detail"]


# ================================================================ the demo path, end to end
def test_kalaimagal_demo_path_end_to_end(client):
    c = client
    sid, r = kalaimagal(c, "demo")
    assert r["status"] == "needs_review" and r["reply"] == M.t("checking", "ta")
    s = view(c, sid)
    top = s["validation"]["likely"][0]
    assert (top["field"], top["suggest"], top["balances"]) == ("weekly.interest_repaid.W3", 310, 3)
    c.post(f"/api/submissions/{sid}/fields/weekly.interest_repaid.W3", json={"value": 310})
    s = c.post(f"/api/submissions/{sid}/fields/weekly.savings.W2", json={"value": 2800}).json()
    assert not errors(s) and s["actions"]["approve"]
    r = c.post(f"/api/submissions/{sid}/approve", json={"officer": "anu"}).json()
    assert r["submission"]["status"] == "awaiting_member" and r["delivered"] and r["channel"] == "simulator"
    items = r["submission"]["summary_items"]
    assert len(items) == 10 and items[4]["value"] == 11400 and items[6]["value"] == 1160
    r = send(c, "demo", text="OK").json()
    assert r["status"] == "written" and r["reply"].startswith("✅")
    assert "95,400" in r["reply"]
    s = view(c, sid)
    assert (s["write_report"]["sheet"], s["write_report"]["column"], len(s["write_report"]["writes"])) == \
        ("Kalaimagal", "G", 16)
    ws = openpyxl.load_workbook(io.BytesIO(c.get("/api/workbook").content))["Kalaimagal"]
    assert ws["G6"].value == 11400 and ws["G8"].value == 1160
