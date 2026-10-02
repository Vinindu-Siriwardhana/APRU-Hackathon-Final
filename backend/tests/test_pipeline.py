"""End to end on synthetic phone photos, with the offline reader standing in for Claude."""
import json
import sys
from pathlib import Path

import cv2
import openpyxl
import pytest

from round1_api_helpers import ROOT, seed  # noqa: E402  (seeds the demo tabs)

from app.extraction import SimulatedReader
from app.imaging.io import read_image
from app.models import FieldStatus
from app.pipeline import Pipeline, Status, Store
from app.schema import load_template

T = load_template()
SYN = ROOT / "samples" / "synthetic"


def photo(name):
    return read_image(SYN / name)


def truth(case):
    return json.loads((SYN / f"{case}.truth.json").read_text(encoding="utf-8"))


@pytest.fixture
def env(tmp_path):
    wb = seed(tmp_path / "gn.xlsx")
    return tmp_path, wb


def make(env, case, errors=None):
    tmp, wb = env
    return Pipeline(Store(tmp / "data"), wb, SimulatedReader(T, truth(case), errors))


def test_clean_report_flows_to_workbook_and_receipt(env):
    p = make(env, "kalaimagal_2026-10")
    sub, reply = p.receive_image("+94770000001", photo("kalaimagal_2026-10_p1.jpg"), lang="ta")
    assert "பக்கம் 2" in reply                                   # asks for page 2, in Tamil
    sub, reply = p.receive_image("+94770000001", photo("kalaimagal_2026-10_p2.jpg"))
    assert sub.status == Status.awaiting_member, [i.message for i in sub.validation.errors]
    assert sub.prior.cash_in_hand == 12500                         # September read from the workbook
    assert "சேமிப்பு" in reply and "11,400" in reply               # Palmera's Tamil label for Savings
    sub, reply = p.member_reply("+94770000001", "சரி")
    assert sub.status == Status.written and "✅" in reply
    ws = openpyxl.load_workbook(env[1])["Kalaimagal"]
    assert ws["G6"].value == 11400 and ws["G2"].value == 4          # June = C, so October = G
    # every value keeps its source location on the photo
    assert sub.record.fields["weekly.savings.W2"].box.page == 1


def test_readings_disagree_goes_to_officer_then_member(env):
    p = make(env, "kalaimagal_2026-10", {"weekly.savings.W2": {"pass": 1, "raw": "2300"}})
    p.receive_image("m1", photo("kalaimagal_2026-10_p1.jpg"))
    sub, reply = p.receive_image("m1", photo("kalaimagal_2026-10_p2.jpg"))
    assert sub.status == Status.needs_review
    assert "weekly.savings.W2" in sub.validation.flagged
    assert (p.store.dir(sub.id) / "crops" / "weekly.savings.W2.jpg").exists()   # officer sees the handwriting
    p.officer_set(sub, "weekly.savings.W2", "2800", officer="officer.anu")
    assert sub.validation.auto_accept
    reply = p.officer_approve(sub, "officer.anu")
    assert sub.status == Status.awaiting_member and "11,400" in reply
    assert sub.record.fields["weekly.savings.W2"].status == FieldStatus.officer_confirmed


def test_confident_misread_is_caught_by_the_ledger(env):
    """Both readings agree on a wrong digit — only the arithmetic can catch it."""
    p = make(env, "kalaimagal_2026-10", {"weekly.savings.W3": {"pass": "both", "raw": "8000"}})
    p.receive_image("m2", photo("kalaimagal_2026-10_p1.jpg"))
    sub, _ = p.receive_image("m2", photo("kalaimagal_2026-10_p2.jpg"))
    assert sub.status == Status.needs_review
    assert {"row_total", "total_income"} <= {i.rule for i in sub.validation.errors}


def test_member_correction_needs_officer(env):
    p = make(env, "sisila_2026-10")
    p.receive_image("m3", photo("sisila_2026-10_p1.jpg"), lang="si")
    sub, reply = p.receive_image("m3", photo("sisila_2026-10_p2.jpg"))
    assert sub.status == Status.awaiting_member, [i.message for i in sub.validation.errors]
    assert "ඉතුරුම්" in reply                                       # Palmera's Sinhala label
    sub, reply = p.member_reply("m3", "5 3600")
    assert sub.status == Status.needs_review and "member_correction" in {i.rule for i in sub.validation.errors}
    p.officer_override(sub, "Savings (Rs.)", accept=False, officer="officer.anu")
    p.officer_approve(sub, "officer.anu")
    sub, reply = p.member_reply("m3", "ok")
    assert sub.status == Status.written


def test_bad_photo_gets_retake_message_in_members_language(env):
    p = make(env, "kalaimagal_2026-10")
    sub, reply = p.receive_image("m4", photo("bad_blurry_p1.jpg"), lang="si")
    assert sub.status == Status.collecting and "පැහැදිලි නැත" in reply
    sub, reply = p.receive_image("m4", photo("bad_glare_p1.jpg"))
    assert "පරාවර්තනයක්" in reply                                   # same conversation stays in Sinhala
    sub, reply = p.receive_image("m5", photo("bad_glare_p1.jpg"), lang="en")
    assert "reflection" in reply


def test_paper_that_does_not_add_up_can_be_accepted_as_written(env):
    """The member's own arithmetic is wrong on paper: once the officer has checked every
    value involved, the check stops blocking but stays visible."""
    p = make(env, "kalaimagal_2026-10", {"weekly.total_income.W1": {"pass": "both", "raw": "6050"}})
    p.receive_image("m6", photo("kalaimagal_2026-10_p1.jpg"))
    sub, _ = p.receive_image("m6", photo("kalaimagal_2026-10_p2.jpg"))
    blocking = {tuple(i.fields) for i in sub.validation.errors}
    assert blocking
    for fields in list(blocking):
        for f in fields:
            if f in sub.record.fields:
                p.officer_set(sub, f, sub.record.fields[f].value, officer="officer.anu")
    assert sub.validation.auto_accept
    assert any(i.message.startswith("Checked against the photo") for i in sub.validation.issues)
