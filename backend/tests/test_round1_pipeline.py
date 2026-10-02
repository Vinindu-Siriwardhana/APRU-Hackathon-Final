"""Round 1 regressions for the pipeline / state machine / bot messages (owner: api).

Includes the pipeline tests ported from the backend review (review/backend/
test_review_regressions.py): superseded reports, editing a written report, a header fix
refreshing the prior month, concurrent writes, a missing workbook row failing the write."""
import json
import threading
import time

import openpyxl
import pytest

from round1_api_helpers import ROOT, seed

from app import messages as M
from app import workbook as W
from app.aggregate import to_monthly
from app.demo_reader import DemoReader, SampleIndex
from app.extraction import ExtractionError, SimulatedReader
from app.imaging.io import read_image
from app.models import FieldStatus, FieldValue, FormRecord, Legibility, PriorMonth
from app.pipeline import NotFound, Pipeline, StateError, Status, Store, Submission
from app.schema import load_template

T = load_template()
SYN = ROOT / "samples" / "synthetic"
CASH = "Cash in Hand – at end of month (mother book) (Rs.)"
KAL_ERRORS = {"weekly.savings.W2": {"pass": 1, "raw": "2300"},
              "weekly.interest_repaid.W3": {"pass": "both", "raw": "810"}}


def photo(name):
    return read_image(SYN / name)


def truth(case):
    return json.loads((SYN / f"{case}.truth.json").read_text(encoding="utf-8"))


@pytest.fixture
def env(tmp_path):
    return tmp_path, seed(tmp_path / "gn.xlsx")


def make(env, case="kalaimagal_2026-10", errors=None, reader=None):
    tmp, wb = env
    return Pipeline(Store(tmp / "data"), wb, reader or SimulatedReader(T, truth(case), errors))


def send_report(p, sender, case="kalaimagal_2026-10", lang="en"):
    p.receive_image(sender, photo(f"{case}_p1.jpg"), lang=lang)
    return p.receive_image(sender, photo(f"{case}_p2.jpg"), lang=lang)


def record(values):
    rec = FormRecord(template_id=T.id)
    for fid, v in values.items():
        rec.fields[fid] = FieldValue(value=v, raw=None if v is None else str(v), passes=[v, v],
                                     legibility=Legibility.blank if v is None else Legibility.clear)
    return rec


def month(**over):
    d = {"header.shg_name": "Kalaimagal", "header.total_members": 15,
         "header.village_gn": "Puthukkudiyiruppu East", "header.month_year": "2026-10"}
    cash = 12500
    for wk, (sav, prin, intr, loans) in zip(["W1", "W2", "W3", "W4"],
                                            [(3000, 2000, 300, 4000), (2800, 2000, 250, 0),
                                             (3000, 2400, 310, 5000), (2600, 2000, 200, 0)]):
        cash += sav + prin + intr - loans
        d.update({f"weekly.meeting_held.{wk}": True, f"weekly.attendance.{wk}": 13,
                  f"weekly.savings.{wk}": sav, f"weekly.principal_repaid.{wk}": prin,
                  f"weekly.interest_repaid.{wk}": intr, f"weekly.loans_distributed.{wk}": loans or None,
                  f"weekly.loan_purpose_income.{wk}": loans or None,
                  f"weekly.total_income.{wk}": sav + prin + intr, f"weekly.total_expenses.{wk}": loans,
                  f"weekly.expected_balance.{wk}": cash, f"weekly.cash_in_hand.{wk}": cash})
    d.update(over)
    return d


@pytest.fixture
def pipe(tmp_path):
    return Pipeline(Store(tmp_path / "data"), seed(tmp_path / "gn.xlsx"), reader=None)


def rules(res, sev="error"):
    return {i.rule for i in res.issues if i.severity.value == sev}


# ================================================================ ported from the backend review
def test_new_report_supersedes_awaiting_one(pipe):
    a = Submission(sender="+94", record=record(month()), status=Status.awaiting_member); pipe.store.save(a)
    time.sleep(0.01)
    b = Submission(sender="+94", record=record(month()), status=Status.needs_review); pipe.store.save(b)
    sub, reply = pipe.member_reply("+94", "OK")
    assert sub is None or sub.id != a.id
    assert pipe.store.load(a.id).status == Status.awaiting_member     # the OK never reached A


def test_officer_cannot_edit_written_submission(pipe):
    sub = Submission(sender="+94", record=record(month()), status=Status.awaiting_member)
    pipe._lookup_prior(sub); pipe.revalidate(sub); pipe.store.save(sub)
    sub, _ = pipe.member_reply("+94", "OK")
    assert sub.status == Status.written
    with pytest.raises(StateError):
        pipe.officer_set(sub, "weekly.savings.W1", "3100", "anu")
    with pytest.raises(StateError):
        pipe.officer_approve(sub, "anu")


def test_header_fix_refreshes_prior(pipe):
    sub = Submission(sender="x", record=record(month(**{"header.month_year": "2026-08"})), status=Status.needs_review)
    pipe._lookup_prior(sub); pipe.revalidate(sub)
    assert not sub.validation.auto_accept                          # August is already recorded
    pipe.officer_set(sub, "header.month_year", "10/2026", "anu")
    assert sub.prior.month == "2026-09" and sub.validation.auto_accept


def test_concurrent_writes_are_not_lost(pipe, monkeypatch):
    orig = W.GNWorkbook.save
    monkeypatch.setattr(W.GNWorkbook, "save", lambda self, path=None: (time.sleep(0.3), orig(self, path))[1])
    # two different groups confirming at the same moment: each from its own sample form
    subs = [Submission(sender="a", record=record(month())),
            Submission(sender="b", record=record(truth("vasantham_2026-10")))]
    ts = [threading.Thread(target=pipe.write, args=(s,)) for s in subs]
    [t.start() for t in ts]; [t.join() for t in ts]
    wb = openpyxl.load_workbook(pipe.workbook_path)
    assert wb["Kalaimagal"]["G6"].value == 11400 and wb["Vasantham"]["G6"].value == 12400
    assert all(s.status == Status.written for s in subs)


def test_missing_required_label_fails_the_write_visibly(pipe):
    wb = openpyxl.load_workbook(pipe.workbook_path); wb["Mapping"]["C16"] = "Loan write off (Rs.)"
    wb.save(pipe.workbook_path)
    sub = Submission(sender="+94", record=record(month()), status=Status.awaiting_member); pipe.store.save(sub)
    sub, reply = pipe.member_reply("+94", "OK")
    assert sub.status == Status.needs_review
    assert "write_failed" in rules(sub.validation)
    assert reply == M.t("confirmed_pending", "en")


# ================================================================ store
def test_store_rejects_bad_ids_and_never_creates_folders_on_read(tmp_path):
    store = Store(tmp_path)
    for bad in ("../../etc", "ABCDEF0123", "x", "0123456789a", "..\\..\\x"):
        with pytest.raises(NotFound):
            store.load(bad)
    with pytest.raises(NotFound):
        store.load("0123456789")
    assert list((tmp_path / "submissions").iterdir()) == []


def test_store_skips_a_torn_file(tmp_path):
    store = Store(tmp_path)
    good = Submission(sender="a"); store.save(good)
    (tmp_path / "submissions" / "0000000000").mkdir()
    (tmp_path / "submissions" / "0000000000" / "submission.json").write_text('{"id": "0000', encoding="utf-8")
    assert [s.id for s in store.all()] == [good.id]
    assert store.open_for("a") is None                            # no crash for any sender


# ================================================================ the Kalaimagal demo path
def test_kalaimagal_demo_path_and_field_statuses(env):
    p = make(env, errors=KAL_ERRORS)
    sub, reply = send_report(p, "m1", lang="ta")
    assert sub.status == Status.needs_review
    assert sub.validation.likely[0].field == "weekly.interest_repaid.W3"
    assert sub.validation.likely[0].suggest == 310
    p.officer_set(sub, "weekly.interest_repaid.W3", "310", "anu")
    p.officer_set(sub, "weekly.savings.W2", 2800, "anu")
    assert sub.validation.auto_accept
    reply = p.officer_approve(sub, "anu")
    assert sub.status == Status.awaiting_member
    assert "அக்டோபர் 2026" in reply and "(ரூ.)" in reply and "(Rs.)" not in reply
    assert reply.en.startswith("Here is what we read") and "October 2026" in reply.en
    bot = [e for e in sub.log if e.get("from") == "bot"][-1]
    assert bot["text_en"] == reply.en
    sub, reply = p.member_reply("m1", "சரி")
    assert sub.status == Status.written
    assert "ரூ. 11,400" in reply and "அக்டோபர்" in reply and "95,400" in reply   # savings to date
    assert "October 2026" in reply.en
    assert not [f for f, fv in sub.record.fields.items() if fv.status == FieldStatus.needs_review]
    assert sub.record.fields["weekly.savings.W1"].status == FieldStatus.member_confirmed
    ws = openpyxl.load_workbook(env[1])["Kalaimagal"]
    assert ws["G6"].value == 11400 and ws["G8"].value == 1160


def test_field_back_to_auto_when_no_longer_flagged(pipe):
    sub = Submission(sender="x", record=record(month(**{"weekly.interest_repaid.W3": 810})), status=Status.needs_review)
    pipe._lookup_prior(sub); pipe.revalidate(sub)
    assert sub.record.fields["weekly.savings.W3"].status == FieldStatus.needs_review
    pipe.officer_set(sub, "weekly.interest_repaid.W3", 310, "anu")
    assert sub.validation.auto_accept
    assert sub.record.fields["weekly.savings.W3"].status == FieldStatus.auto


# ================================================================ guards
def test_status_guards(env):
    p = make(env, errors=KAL_ERRORS)
    sub, _ = p.receive_image("g1", photo("kalaimagal_2026-10_p1.jpg"))
    with pytest.raises(StateError):                               # collecting: nothing read yet
        p.officer_set(sub, "weekly.savings.W1", 1, "anu")
    with pytest.raises(StateError):
        p.officer_approve(sub, "anu")
    sub, _ = p.receive_image("g1", photo("kalaimagal_2026-10_p2.jpg"))
    with pytest.raises(StateError, match="still failing"):
        p.officer_approve(sub, "anu")
    with pytest.raises(ValueError):                              # force needs a reason
        p.officer_approve(sub, "anu", force=True)
    with pytest.raises(NotFound):
        p.officer_override(sub, "Savings (Rs.)", True, "anu")
    with pytest.raises(NotFound):
        p.officer_set(sub, "weekly.nonsense.W1", 1, "anu")
    for bad in (True, float("nan"), float("inf"), -5, 1e12, [1, 2], "abc"):
        with pytest.raises(ValueError):
            p.officer_set(sub, "weekly.savings.W1", bad, "anu")
    with pytest.raises(ValueError):
        p.officer_set(sub, "weekly.attendance.W1", 1.5, "anu")
    p.officer_approve(sub, "anu", force=True, reason="checked with the group by phone")
    assert sub.status == Status.awaiting_member and sub.forced["reason"]
    with pytest.raises(StateError, match="waiting for her reply"):    # no double summary
        p.officer_approve(sub, "anu")


def test_editing_an_awaiting_report_sends_it_back_to_review(env):
    p = make(env)
    sub, _ = send_report(p, "e1")
    assert sub.status == Status.awaiting_member
    p.officer_set(sub, "weekly.savings.W1", 3100, "anu")
    assert sub.status == Status.needs_review and not sub.member_confirmed


# ================================================================ superseding + routing
def test_new_photo_supersedes_open_report_for_same_group_and_month(env):
    p = make(env, errors=KAL_ERRORS)
    a, _ = send_report(p, "s1")
    assert a.status == Status.needs_review
    b, _ = send_report(p, "s2")                                  # another phone, same SHG + month
    a = p.store.load(a.id)
    assert a.status == Status.superseded and a.superseded_by == b.id
    sub, reply = p.member_reply("s1", "OK")                       # her old report takes no replies
    assert sub is None and reply == M.t("no_report", "en")


def test_text_goes_to_newest_report_and_answers_by_state(env):
    p = make(env, errors=KAL_ERRORS)
    sub, _ = send_report(p, "r1", lang="ta")
    _, reply = p.member_reply("r1", "OK")                         # still in review: told so, in Tamil
    assert reply == M.t("officer_checking", "ta") and reply.en == M.t("officer_checking", "en")
    p.receive_image("r1", photo("bad_blurry_p1.jpg"))              # a bad photo doesn't steal her replies
    sub2, reply = p.member_reply("r1", "OK")
    assert sub2.id == sub.id
    p.receive_image("r1", photo("kalaimagal_2026-10_p1.jpg"))     # a new report has started
    sub3, reply = p.member_reply("r1", "hello")
    assert sub3.id != sub.id and reply == M.t("waiting_page2", "ta")


# ================================================================ member corrections
def test_member_correction_of_a_sum_needs_the_weekly_cells(env):
    p = make(env)
    sub, _ = send_report(p, "c1")
    sub, reply = p.member_reply("c1", "5 Rs.12,000/=")
    assert "changed" not in reply and "noted" in reply
    assert sub.status == Status.needs_review
    issue = next(i for i in sub.validation.errors if i.rule == "member_correction")
    assert issue.expected == 12000 and issue.found == 11400
    assert "weekly.savings.W1" in issue.fields and "weekly.savings.W4" in issue.fields
    with pytest.raises(StateError, match="weekly cells"):
        p.officer_override(sub, "Savings (Rs.)", True, "anu")
    p.officer_set(sub, "weekly.savings.W1", sub.record.get("weekly.savings.W1") + 600, "anu")
    assert "member_correction" not in rules(sub.validation)
    assert any(e.get("event") == "member_correction_resolved" for e in sub.log)
    assert sub.overrides == {}
    assert not sub.validation.auto_accept                         # the ledger now disagrees: officer must look


def test_member_correction_of_a_single_cell_can_be_accepted(env):
    p = make(env)
    sub, _ = send_report(p, "c2")
    sub, _ = p.member_reply("c2", "9 - 20,000")
    corr = p.corrections(sub)[0]
    assert corr["can_accept"] and corr["cell"].startswith("weekly.cash_in_hand.")
    p.officer_override(sub, CASH, True, "anu")
    assert to_monthly(sub.record, T)[CASH] == 20000
    assert sub.record.fields[corr["cell"]].status == FieldStatus.member_corrected
    # re-checked: her figure doesn't match the ledger, so it blocks until the officer looks
    assert not sub.validation.auto_accept
    assert any(corr["cell"] in i.fields for i in sub.validation.errors)


def test_member_correction_dismissed_keeps_the_form(env):
    p = make(env, "sisila_2026-10")
    sub, _ = send_report(p, "c3", "sisila_2026-10", lang="si")
    before = to_monthly(sub.record, T)["Savings (Rs.)"]
    sub, reply = p.member_reply("c3", "5 3600")
    assert reply == M.t("correction_noted", "si", n=5, value="3,600")
    p.officer_override(sub, "Savings (Rs.)", accept=False, officer="anu")
    p.officer_approve(sub, "anu")
    sub, _ = p.member_reply("c3", "හරි")
    assert sub.status == Status.written
    assert sub.write_report.writes and to_monthly(sub.record, T)["Savings (Rs.)"] == before


def test_member_sends_the_same_figure(env):
    p = make(env)
    sub, _ = send_report(p, "c4")
    sub, reply = p.member_reply("c4", "5 11400")
    assert sub.status == Status.awaiting_member and "already says 11,400" in reply


def test_ambiguous_reply_is_not_guessed(env):
    p = make(env)
    sub, _ = send_report(p, "c5")
    for text in ("300 310", "5 300 310", "OK 5 12000"):
        sub, reply = p.member_reply("c5", text)
        assert sub.status == Status.awaiting_member and sub.overrides == {}
        assert "5 12000" in reply


# ================================================================ writing
def test_write_failure_is_never_silent_and_officer_can_retry(env, monkeypatch):
    p = make(env)
    sub, _ = send_report(p, "w1", lang="ta")

    def locked(self, path=None):
        raise PermissionError("locked")
    monkeypatch.setattr(W.GNWorkbook, "save", locked)
    sub, reply = p.member_reply("w1", "OK")
    assert sub.status == Status.needs_review and reply == M.t("confirmed_pending", "ta")
    issue = next(i for i in sub.validation.errors if i.rule == "write_failed")
    assert "Close" in issue.message and "Excel" in issue.message
    monkeypatch.undo()
    reply = p.officer_approve(sub, "anu")                        # she already said OK: just write
    assert sub.status == Status.written and "✅" in reply
    assert "write_failed" not in rules(sub.validation)


def test_duplicate_month_is_caught_before_the_member_confirms(env):
    p = make(env)
    a, _ = send_report(p, "d1")
    p.member_reply("d1", "OK")
    b, _ = send_report(p, "d2")
    assert b.status == Status.needs_review and "month_already_recorded" in rules(b.validation)


def test_unknown_group_blocks_until_officer_confirms_new_group(pipe):
    sub = Submission(sender="n1", record=record(month(**{"header.shg_name": "Nilam"})), status=Status.needs_review)
    pipe._lookup_prior(sub); pipe.revalidate(sub)
    assert "unknown_shg" in rules(sub.validation)
    pipe.confirm_new_group(sub, "anu")
    assert "unknown_shg" in rules(sub.validation, "warning") and sub.validation.auto_accept
    pipe.officer_approve(sub, "anu"); pipe.store.save(sub)
    sub, _ = pipe.member_reply("n1", "OK")
    assert sub.status == Status.written and sub.write_report.created_tab


# ================================================================ reading failures + demo reader
class Broken:
    def read(self, images, prompt, schema):
        raise ExtractionError("The reading service could not be reached.")


def test_extraction_failure_then_retry(env):
    p = make(env, reader=Broken())
    sub, reply = send_report(p, "f1", lang="si")
    assert sub.status == Status.failed and "could not be reached" in sub.failure
    assert reply == M.t("extraction_failed", "si")
    assert any(e.get("event") == "extraction_failed" for e in sub.log)
    _, reply = p.member_reply("f1", "OK")
    assert reply == M.t("officer_checking", "si")
    p.reader = SimulatedReader(T, truth("kalaimagal_2026-10"))
    reply = p.retry(sub.id, "anu")
    assert p.store.load(sub.id).status == Status.awaiting_member and "11,400" in reply


def test_demo_reader_does_not_invent_data(env):
    p = make(env, reader=DemoReader(T, SampleIndex(SYN)))
    p.receive_image("u1", photo("kalaimagal_2026-10_p1.jpg"), source=None)   # a re-saved / unknown photo
    sub, reply = p.receive_image("u1", photo("kalaimagal_2026-10_p2.jpg"), source="kalaimagal_2026-10_p2.jpg")
    assert sub.status == Status.failed and sub.record is None
    assert reply == M.t("demo_unknown_photo", "en")
    p.receive_image("u2", photo("kalaimagal_2026-10_p1.jpg"), source="kalaimagal_2026-10_p1.jpg")
    sub, _ = p.receive_image("u2", photo("vasantham_2026-10_p2.jpg"), source="vasantham_2026-10_p2.jpg")
    assert sub.status == Status.failed and "different sample reports" in sub.failure


def test_process_now_reads_page1_alone(env):
    p = make(env)
    sub, _ = p.receive_image("p1", photo("kalaimagal_2026-10_p1.jpg"))
    p.process_now(sub.id, "anu")
    sub = p.store.load(sub.id)
    assert sub.status in (Status.needs_review, Status.awaiting_member) and sub.record is not None
    with pytest.raises(StateError):
        p.process_now(sub.id, "anu")


def test_reject_tells_the_member_why(env):
    p = make(env, errors=KAL_ERRORS)
    sub, _ = send_report(p, "j1", lang="ta")
    reply = p.officer_reject(sub, "anu", "unclear", note="too dark at the bottom")
    assert sub.status == Status.rejected
    assert M.REJECT_REASONS["unclear"]["ta"] in reply and "too dark" not in reply
    assert M.REJECT_REASONS["unclear"]["en"] in reply.en
    with pytest.raises(StateError):
        p.officer_reject(sub, "anu", "unclear")


# ================================================================ front door: privacy, STOP, language
def test_privacy_note_once_and_stop(env):
    p = make(env)
    _, r1 = p.on_text("+9471", "hello")
    assert M.t("privacy", "en") in r1
    _, r2 = p.on_text("+9471", "hello")
    assert M.t("privacy", "en") not in r2
    p.on_image("+9471", photo("kalaimagal_2026-10_p1.jpg"))
    _, r3 = p.on_text("+9471", "STOP")
    assert r3 == M.t("stopped", "en")
    assert all(s.status == Status.rejected for s in p.store.for_sender("+9471"))
    conv = p.conversations.load("+9471")
    assert conv.opted_out and any(e["event"] == "opted_out" for e in conv.events)


def test_language_from_keyword_or_script(env):
    p = make(env)
    _, r = p.on_text("+9472", "தமிழ்")
    assert r.startswith(M.t("lang_set", "ta"))
    _, r = p.on_text("+9473", "හරි")                               # writes in Sinhala: replies in Sinhala
    assert r.startswith(M.t("no_report", "si"))
    assert p.conversations.load("+9473").lang == "si"


def test_conversation_log_rebuilds_the_phone(env):
    p = make(env)
    p.on_image("sim:x", photo("kalaimagal_2026-10_p1.jpg"), lang="ta", image_name="kalaimagal_2026-10_p1.jpg")
    p.on_text("sim:x", "OK")
    msgs = p.conversations.load("sim:x").messages
    assert [m["from"] for m in msgs] == ["me", "bot", "me", "bot"]
    assert msgs[0]["image"] == "kalaimagal_2026-10_p1.jpg" and msgs[1]["text_en"]


# ================================================================ messages
@pytest.mark.parametrize("text,n,value", [
    ("5 12000", 5, 12000), ("5 12,000", 5, 12000), ("5. 12000", 5, 12000), ("5 Rs.12000/=", 5, 12000),
    ("5 12,000/-", 5, 12000), ("5 - 12000", 5, 12000), ("5: 12000", 5, 12000), ("௫ ௧௨௦௦௦", 5, 12000),
    ("9 රු. 20,160", 9, 20160), ("5 1,20,000", 5, 120000),
])
def test_member_reply_formats(text, n, value):
    kind, (item, raw) = M.parse_member_reply(text)
    assert kind == "correct" and item == n and M.parse_amount(raw) == value


@pytest.mark.parametrize("text", ["300 310", "512000", "10 5000", "hello", "OK 5 12000", ""])
def test_member_reply_rejects_ambiguous(text):
    kind, arg = M.parse_member_reply(text)
    if kind == "correct":                                        # e.g. never for these
        with pytest.raises(ValueError):
            M.parse_amount(arg[1])
    else:
        assert kind == "unknown"


@pytest.mark.parametrize("raw", ["300 310", "12.000.5", "1e5", "12000abc", "-500"])
def test_parse_amount_refuses(raw):
    with pytest.raises(ValueError):
        M.parse_amount(raw)


@pytest.mark.parametrize("text", ["OK", "ok 👍", "Ok thank you", "சரி", "හරි", "ஆம்", "✅", "yes, all correct"])
def test_confirm_words(text):
    assert M.parse_member_reply(text)[0] == "confirm"


def test_receipt_uses_month_name_and_dash_for_missing():
    r = M.receipt("si", month="2026-10", group="Sisila", savings=4000, repay=2100, cash=None, savings_to_date=44000)
    assert "ඔක්තෝබර් 2026" in r and "රු. 4,000" in r and ": —" in r and "රු. 44,000" in r
    assert "October 2026" in r.en and "Rs. 44,000" in r.en and "Rs 0" not in r.en


def test_correction_noted_does_not_claim_a_change():
    for lang in M.LANGS:
        assert "changed" not in M.t("correction_noted", "en", n=5, value="12,000")
    assert "noted" in M.t("correction_noted", "en", n=5, value="12,000")


def test_every_message_has_all_three_languages():
    for key, d in {**M.TEXT, **M.REJECT_REASONS}.items():
        assert set(M.LANGS) <= set(d), key
