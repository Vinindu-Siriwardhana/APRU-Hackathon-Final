"""Write verified monthly figures into Palmera's GN-level workbook.

How Palmera's GN workbook works (from the files they shared):
* one tab per SHG, named after the SHG; months run across columns C..Z, starting
  from the date in C1 (column C = month 1 and also holds the opening "to date" totals)
* rows are labelled in column B via =Mapping!Cn; totals, "cash from app" and all
  "to date" rows (after month 1) are formulas
* 'SHG Monitoring' has a 29-row block per SHG that looks values up by label, and the
  district workbook pulls that sheet in with IMPORTRANGE

So the only cells anyone ever types into are the input rows of one month column on
one SHG tab. That is all this module writes. It never overwrites a formula, it
refuses to change an existing value unless an officer confirmed a correction
(overwrite=True), it only creates a tab for a group an officer confirmed is new
(create_tab=True), and it saves atomically (temp file + rename) so a crash or a
second writer can never leave a half-written .xlsx behind.

The same interface can be backed by the Google Sheets API for the live sheet; this
implementation works on an .xlsx export (openpyxl) for testing and offline demos.
"""
from __future__ import annotations

import os
import re
import shutil
import tempfile
import unicodedata
from copy import copy
from datetime import datetime
from pathlib import Path
from typing import Any, Optional

import openpyxl
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter
from openpyxl.workbook.views import BookView
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import BaseModel, Field

from .models import PriorMonth, parse_number
from .schema import Template

TEMPLATE_TAB = "<name of SHG>"
MONITORING_TAB = "SHG Monitoring"
MONITORING_BLOCK = (2, 30)          # rows of the template SHG's block in 'SHG Monitoring'
FIRST_COL, LAST_COL = 3, 26         # C..Z


class WorkbookError(Exception):
    pass


class CellWrite(BaseModel):
    label: str
    cell: str
    old: Any = None
    new: Any = None


class WriteReport(BaseModel):
    sheet: str
    month: str
    column: str
    created_tab: bool = False
    writes: list[CellWrite] = Field(default_factory=list)
    skipped: list[str] = Field(default_factory=list)       # labels with no row in this workbook
    warnings: list[str] = Field(default_factory=list)


_FORBIDDEN = re.compile(r"[\[\]:*?/\\]")


def tab_title(name: str) -> str:
    """The tab name Excel allows for an SHG name: no []:*?/\\, at most 31 characters, and
    no straight apostrophe (Palmera's INDIRECT("'"&name&"'!…") formulas would break)."""
    t = _FORBIDDEN.sub("-", unicodedata.normalize("NFC", str(name)).strip()).replace("'", "’")
    return t[:31].strip()


def _norm(s: str) -> str:
    s = unicodedata.normalize("NFKC", str(s)).lower()
    s = s.replace("–", "-").replace("—", "-")
    return re.sub(r"[^a-z0-9඀-෿஀-௿]+", " ", s).strip()


def _month_key(d: Any) -> str:
    if isinstance(d, datetime):
        return f"{d.year:04d}-{d.month:02d}"
    raise WorkbookError(f"C1 is not a date: {d!r}")


def _add_months(ym: str, n: int) -> str:
    y, m = map(int, ym.split("-"))
    idx = y * 12 + (m - 1) + n
    return f"{idx // 12:04d}-{idx % 12 + 1:02d}"


def _months_between(a: str, b: str) -> int:
    ya, ma = map(int, a.split("-"))
    yb, mb = map(int, b.split("-"))
    return (yb * 12 + mb) - (ya * 12 + ma)


class GNWorkbook:
    def __init__(self, path: str | Path, template: Template):
        self.path = Path(path)
        self.t = template
        self.wb = openpyxl.load_workbook(self.path)
        self.mapping = self.wb["Mapping"] if "Mapping" in self.wb.sheetnames else None
        self._written: Optional[str] = None          # tab to leave active on save

    # ------------------------------------------------------------------ lookup
    def shg_tabs(self) -> list[str]:
        fixed = {"Mapping", MONITORING_TAB, "Functioning SHG details", TEMPLATE_TAB}
        return [n for n in self.wb.sheetnames if n not in fixed]

    def find_shg(self, name: str) -> tuple[Optional[str], list[str]]:
        """Exact (normalised) match, else candidates for the officer to pick from."""
        target = _norm(name)
        as_tab = _norm(tab_title(name))              # long names are cut to 31 characters
        tabs = self.shg_tabs()
        for n in tabs:
            if _norm(n) in (target, as_tab):
                return n, []
        from difflib import get_close_matches
        normed = {_norm(n): n for n in tabs}
        close = get_close_matches(target, list(normed), n=3, cutoff=0.6)
        return None, [normed[c] for c in close]

    def _label(self, ws: Worksheet, row: int) -> str:
        v = ws.cell(row, 2).value
        if isinstance(v, str) and v.startswith("=") and self.mapping is not None:
            m = re.fullmatch(r"=Mapping!\$?([A-Z]+)\$?(\d+)", v.strip())
            if m:
                v = self.mapping[f"{m[1]}{m[2]}"].value
        return "" if v is None else str(v)

    def label_rows(self, ws: Worksheet) -> dict[str, int]:
        return {_norm(self._label(ws, r)): r for r in range(2, ws.max_row + 1) if self._label(ws, r)}

    def start_month(self, ws: Worksheet) -> str:
        return _month_key(ws["C1"].value)

    def month_col(self, ws: Worksheet, month: str) -> int:
        start = self.start_month(ws)
        col = FIRST_COL + _months_between(start, month)
        if not FIRST_COL <= col <= LAST_COL:
            raise WorkbookError(f"{month} is outside this tab's range "
                                f"({start} to {_add_months(start, LAST_COL - FIRST_COL)}).")
        return col

    # ------------------------------------------------------------------ reading
    def _num(self, ws: Worksheet, row: Optional[int], col: int) -> Optional[float]:
        """A typed input value. Blank, a dash or a formula -> None. Other text raises a
        WorkbookError naming the cell (never a bare ValueError that crashes the caller)."""
        if row is None:
            return None
        cell = ws.cell(row, col)
        v = cell.value
        if v is None or isinstance(v, (bool, ArrayFormula)):
            return None
        if isinstance(v, (int, float)):
            return float(v)
        if isinstance(v, str):
            if v.strip().startswith("="):
                return None
            try:
                n = parse_number(v)
            except ValueError:
                n = "bad"
            if n != "bad":
                return None if n is None else float(n)
        raise WorkbookError(f"{ws.title}!{cell.coordinate} holds {v!r}, which is not a number. "
                            "Please fix that cell in the workbook.")

    def _sum(self, ws: Worksheet, row: Optional[int], c0: int, c1: int) -> float:
        return sum(self._num(ws, row, c) or 0 for c in range(c0, c1 + 1))

    def _input_rows(self, rows: dict[str, int]) -> list[int]:
        """Rows of the 16 monthly input figures that exist on this tab."""
        out = []
        for spec in self.t.workbook["monthly_rows"]:
            r = rows.get(_norm(spec["label"]))
            if r is not None:
                out.append(r)
        return out

    def _col_filled(self, ws: Worksheet, rows: dict[str, int], col: int) -> bool:
        return any(ws.cell(r, col).value not in (None, "") for r in self._input_rows(rows))

    def month_recorded(self, sheet: str, month: str) -> bool:
        """True if any of the monthly input cells for `month` already holds something."""
        ws = self.wb[sheet]
        return self._col_filled(ws, self.label_rows(ws), self.month_col(ws, month))

    def prior_month(self, sheet: str, month: str) -> PriorMonth:
        """Month-end balances for the last month recorded BEFORE `month`, recomputed from
        inputs (openpyxl can't evaluate formulas; the Sheets API backend reads them directly).
        Loans outstanding uses the corrected ledger, not the sheet's current formula.
        `month` is None when nothing is recorded before; validation only uses the balances
        when `month` is exactly the month before (otherwise it warns about the gap)."""
        ws = self.wb[sheet]
        rows = self.label_rows(ws)
        r = lambda label: rows.get(_norm(label))
        this_col = self.month_col(ws, month)
        recorded = self._col_filled(ws, rows, this_col)
        col = this_col - 1
        while col >= FIRST_COL and not self._col_filled(ws, rows, col):
            col -= 1                                   # walk back over months not entered
        if col < FIRST_COL:
            return PriorMonth(this_month_recorded=recorded)
        sav = (self._num(ws, r("Total savings to date (Rs.)"), FIRST_COL) or 0) + \
            self._sum(ws, r("Savings (Rs.)"), FIRST_COL + 1, col)
        refunded = (self._num(ws, r("Total Savings refunded to date (Rs.)"), FIRST_COL) or 0) + \
            self._sum(ws, r("Savings refunded (Rs.)"), FIRST_COL + 1, col)
        out = (self._num(ws, r("Total Loans outstanding to date (Rs.)"), FIRST_COL) or 0) + \
            self._sum(ws, r("Loans distributed (Rs.)"), FIRST_COL + 1, col) - \
            self._sum(ws, r("Principal loan repayments (Rs.)"), FIRST_COL + 1, col) - \
            self._sum(ws, r("Loan right off (Rs.)"), FIRST_COL + 1, col)
        members = self._num(ws, r("Number of SHG members"), col)
        return PriorMonth(
            month=_add_months(self.start_month(ws), col - FIRST_COL),
            cash_in_hand=self._num(ws, r("Cash in Hand – at end of month (mother book) (Rs.)"), col),
            savings_to_date=sav, savings_refunded_to_date=refunded, loans_outstanding=out,
            members=int(members) if members is not None else None,
            this_month_recorded=recorded,
        )

    def gn_of(self, sheet: str) -> Optional[str]:
        mon = self.wb[MONITORING_TAB]
        for r in range(2, mon.max_row + 1):
            if mon.cell(r, 2).value == sheet:
                return mon.cell(r, 1).value
        return None

    def series(self, sheet: str) -> list[dict[str, Any]]:
        """Month-by-month figures for one SHG, with ledger balances recomputed here
        (loans outstanding with the corrected formula; cash without counting write-offs)."""
        ws = self.wb[sheet]
        rows = self.label_rows(ws)
        r = lambda label: rows.get(_norm(label))
        start = self.start_month(ws)
        out: list[dict[str, Any]] = []
        sav = self._num(ws, r("Total savings to date (Rs.)"), FIRST_COL) or 0
        outst = self._num(ws, r("Total Loans outstanding to date (Rs.)"), FIRST_COL) or 0
        opening = [self._num(ws, r(l), FIRST_COL) or 0 for l in (
            "Total savings to date (Rs.)", "Total Principal loan repayments to date (Rs.)",
            "Total Interest repayments to date (Rs.)", "Total Other income to date (Rs.)")]
        spent = [self._num(ws, r(l), FIRST_COL) or 0 for l in (
            "Total Loans distributed to date (Rs.)", "Total Savings refunded to date (Rs.)",
            "Total Other expenses to date (Rs.)")]
        ledger_cash = sum(opening) - sum(spent)
        for col in range(FIRST_COL, LAST_COL + 1):
            g = lambda label: self._num(ws, r(label), col)
            if not self._col_filled(ws, rows, col):
                break
            meetings = g("Total number of meetings held") or 0
            members = g("Number of SHG members") or 0
            att = g("Total attendance for the month") or 0
            m = {
                "month": _add_months(start, col - FIRST_COL), "column": get_column_letter(col),
                "members": members, "meetings": meetings, "attendance": att,
                "attendance_rate": round(att / (members * meetings), 3) if members and meetings else None,
                "savings": g("Savings (Rs.)") or 0, "principal": g("Principal loan repayments (Rs.)") or 0,
                "interest": g("Interest repayment (Rs.)") or 0,
                "other_income": g("Other income (Membership fees, Other NGO funds, Fines, Social fund, Other) (Rs.)") or 0,
                "loans": g("Loans distributed (Rs.)") or 0, "refunds": g("Savings refunded (Rs.)") or 0,
                "other_expenses": g("Other expenses (Interest refunded, Community / members support, CLA fee, Admin expenses & other) (Rs.)") or 0,
                "write_offs": g("Loan right off (Rs.)") or 0,
                "cash": g("Cash in Hand – at end of month (mother book) (Rs.)"),
            }
            if col > FIRST_COL:
                sav += m["savings"]
                outst += m["loans"] - m["principal"] - m["write_offs"]
                ledger_cash += (m["savings"] + m["principal"] + m["interest"] + m["other_income"]
                                - m["loans"] - m["refunds"] - m["other_expenses"])
            m.update(savings_to_date=sav, loans_outstanding=outst, ledger_cash=ledger_cash)
            out.append(m)
        return out

    # ------------------------------------------------------------------ writing
    def create_shg_tab(self, shg_name: str, gn_name: str, start_month: str) -> Worksheet:
        """Copy the blank '<name of SHG>' tab and register the SHG in 'SHG Monitoring'."""
        title = tab_title(shg_name)
        if not title:
            raise WorkbookError("The SHG name is empty, so no tab can be created.")
        if title.lower() in (n.lower() for n in self.wb.sheetnames):
            raise WorkbookError(f"Tab '{title}' already exists.")
        src = self.wb[TEMPLATE_TAB]
        ws = self.wb.copy_worksheet(src)
        ws.title = title
        for dv in src.data_validations.dataValidation:           # copy_worksheet drops these…
            ws.add_data_validation(copy(dv))
        for cf in src.conditional_formatting:                     # …and Palmera's red highlights
            for rule in cf.rules:
                ws.conditional_formatting.add(str(cf.sqref), copy(rule))
        y, m = map(int, start_month.split("-"))
        ws["C1"] = datetime(y, m, 1)
        ws["C1"].number_format = src["C1"].number_format
        self._register_in_monitoring(ws.title, gn_name)
        return ws

    def _register_in_monitoring(self, shg: str, gn: str) -> None:
        mon = self.wb[MONITORING_TAB]
        a, b = MONITORING_BLOCK
        # first free row after the last used block
        last = max((r for r in range(2, mon.max_row + 1) if mon.cell(r, 2).value), default=1)
        start = last + 1
        for i, src_row in enumerate(range(a, b + 1)):
            dst = start + i
            for c in range(1, mon.max_column + 1):
                s = mon.cell(src_row, c)
                d = mon.cell(dst, c)
                d.value = _translated(s.value, s.coordinate, d.coordinate, dst - src_row)
                if s.has_style:
                    d._style = copy(s._style)
            mon.cell(dst, 1).value = gn
            mon.cell(dst, 2).value = shg

    def write_month(self, shg_name: str, month: str, values: dict[str, Any], *,
                    gn_name: Optional[str] = None, overwrite: bool = False, create_tab: bool = False,
                    opening: Optional[dict[str, Any]] = None, correction: bool = False) -> WriteReport:
        """Write one month's figures. overwrite=True (alias: correction) replaces values
        already there: only after an officer confirmed a correction. create_tab=True creates
        a tab for an unknown name: only after an officer confirmed it is a new group."""
        overwrite = overwrite or correction
        sheet, candidates = self.find_shg(shg_name)
        created = False
        if sheet is None:
            if not create_tab:
                if candidates:
                    raise WorkbookError(f"SHG '{shg_name}' not found. Did you mean: {', '.join(candidates)}?")
                raise WorkbookError(f"No tab for SHG '{shg_name}' — confirm it as a new group first.")
            if not gn_name:
                raise WorkbookError(f"No tab for SHG '{shg_name}', and no Village / GN to register it under.")
            sheet = self.create_shg_tab(shg_name, gn_name, month).title
            created = True
        ws = self.wb[sheet]
        col = self.month_col(ws, month)
        letter = get_column_letter(col)
        rows = self.label_rows(ws)
        rep = WriteReport(sheet=sheet, month=month, column=letter, created_tab=created)

        todo = dict(values)
        if col == FIRST_COL:
            # month 1: column C also carries the opening "to date" totals as plain inputs
            todo.update(self._opening_values(values, opening, rep))
        optional = {_norm(r["label"]) for r in self.t.workbook.get("optional_rows", [])}

        # check first, then write — never a half-written month
        planned: list[tuple[int, str, Any]] = []
        missing: list[str] = []
        for label, new in todo.items():
            row = rows.get(_norm(label))
            if row is None:
                if _norm(label) in optional:
                    rep.skipped.append(label)        # e.g. a row Palmera hasn't added yet
                else:
                    missing.append(label)
                continue
            cell = ws.cell(row, col)
            old = cell.value
            if isinstance(old, ArrayFormula) or (isinstance(old, str) and old.startswith("=")):
                raise WorkbookError(f"{sheet}!{cell.coordinate} ({label}) is a formula; refusing to overwrite.")
            if new is None:
                rep.warnings.append(f"{label}: no figure on the form, so {sheet}!{cell.coordinate} was left as it is.")
                continue
            if old not in (None, "") and old != new and not overwrite:
                raise WorkbookError(f"{sheet}!{cell.coordinate} ({label}) already holds {old!r}. "
                                    f"An officer must confirm a correction to replace it with {new!r}.")
            planned.append((row, label, new))
        if missing:
            raise WorkbookError(f"Tab '{sheet}' has no row labelled {', '.join(repr(m) for m in missing)} "
                                "(column B). The workbook layout has changed, so nothing was written.")
        for row, label, new in planned:
            cell = ws.cell(row, col)
            rep.writes.append(CellWrite(label=label, cell=f"{sheet}!{cell.coordinate}", old=cell.value, new=new))
            cell.value = new
        self._written = sheet
        return rep

    def _opening_values(self, values: dict[str, Any], opening: Optional[dict[str, Any]],
                        rep: WriteReport) -> dict[str, Any]:
        """To-date totals for column C. Given explicitly (from the group's mother book) or,
        for a group whose records start this month, equal to this month's flows."""
        if opening:
            return dict(opening)
        g = lambda k: values.get(k) or 0
        rep.warnings.append("New SHG tab: opening 'to date' totals were set to this month's figures. "
                            "That is only right for a group that started this month — otherwise Palmera's "
                            "'cash from app' will be off by the cash the group already held. Enter the "
                            "lifetime totals from the mother book.")
        return {
            "Total savings to date (Rs.)": g("Savings (Rs.)"),
            "Total Principal loan repayments to date (Rs.)": g("Principal loan repayments (Rs.)"),
            "Total Interest repayments to date (Rs.)": g("Interest repayment (Rs.)"),
            "Total Other income to date (Rs.)": g("Other income (Membership fees, Other NGO funds, Fines, Social fund, Other) (Rs.)"),
            "Total Loans distributed to date (Rs.)": g("Loans distributed (Rs.)"),
            "Total Savings refunded to date (Rs.)": g("Savings refunded (Rs.)"),
            "Total Other expenses to date (Rs.)": g("Other expenses (Interest refunded, Community / members support, CLA fee, Admin expenses & other) (Rs.)"),
            "Total Loan right off to date (Rs.)": g("Loan right off (Rs.)"),
            "Total Loans outstanding to date (Rs.)": g("Loans distributed (Rs.)") - g("Principal loan repayments (Rs.)") - g("Loan right off (Rs.)"),
        }

    def _activate(self, sheet: str) -> None:
        """Open the workbook on the tab just written (and only that tab selected: two
        selected tabs would make Excel group them, so typing would edit both)."""
        ws = self.wb[sheet]
        for other in self.wb.worksheets:
            other.sheet_view.tabSelected = other is ws
        self.wb.active = ws
        if not self.wb.views:                      # Palmera's file has no <bookViews>: without one
            self.wb.views.append(BookView())       # the active tab isn't saved at all

    def save(self, path: str | Path | None = None) -> Path:
        """Atomic: write a temp file in the same folder, then rename it over the original,
        so a crash or a concurrent reader never sees a half-written .xlsx."""
        out = Path(path) if path else self.path
        if self._written and self._written in self.wb.sheetnames:
            self._activate(self._written)
        fd, tmp = tempfile.mkstemp(prefix=f".{out.stem}-", suffix=".xlsx", dir=out.parent)
        os.close(fd)
        try:
            self.wb.save(tmp)
            if out.exists():
                shutil.copymode(out, tmp)
            os.replace(tmp, out)
        except PermissionError as e:             # Windows: Excel has the file open
            raise WorkbookError(f"Close {out.name} in Excel and try again.") from e
        finally:
            if os.path.exists(tmp):
                os.unlink(tmp)
        return out


def _translated(v: Any, origin: str, dest: str, row_shift: int) -> Any:
    """A formula copied from `origin` to `dest`, references moved like Excel's copy/paste.
    Array formulas need their text AND their range moved, or Excel reports the file as damaged."""
    if isinstance(v, ArrayFormula):
        ref = CellRange(v.ref)
        ref.shift(row_shift=row_shift)
        return ArrayFormula(ref=ref.coord, text=Translator(v.text, origin=origin).translate_formula(dest))
    if isinstance(v, str) and v.startswith("="):
        return Translator(v, origin=origin).translate_formula(dest)
    return v
